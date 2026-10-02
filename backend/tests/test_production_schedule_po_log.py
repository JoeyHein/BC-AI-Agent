"""SO-PO Log sheet — the persistent record of which sales order is linked to
which purchase order and when each link was last touched. Unlike the other
read-only sheets it carries history across refreshes. See
production_schedule_service.build_po_log.
"""
import io
import os
import sys
from datetime import datetime

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from openpyxl import load_workbook

from app.services.production_schedule_service import production_schedule_service as svc

T1 = datetime(2026, 10, 2, 9, 0)
T2 = datetime(2026, 10, 2, 12, 30)
T3 = datetime(2026, 10, 3, 4, 30)


def _so(number, customer="Acme Doors"):
    return {"number": number, "customerName": customer, "orderDate": "2026-10-01",
            "salesOrderLines": [{"lineType": "Item", "lineObjectNumber": "PN45-1", "quantity": 1}]}


def _line(po="PO-001", status="Draft", qty=4, received=0):
    return {"lineType": "Item", "lineObjectNumber": "PN45-1", "quantity": qty,
            "receivedQuantity": received, "expectedReceiptDate": "2026-10-20",
            "_po_number": po, "_po_status": status, "_vendor": "UPWARDOR", "_po_date": "2026-10-01"}


def _run(orders, po_lines_by_so, prior_bytes, now):
    """One refresh: read back the prior file, rebuild, return (bytes, log rows)."""
    prior = svc.parse_po_log_from_bytes(prior_bytes) if prior_bytes else {}
    facts = svc.compute_so_facts(orders, po_lines_by_so)
    data, _, _ = svc.build_workbook_bytes(orders, {}, facts, po_lines_by_so=po_lines_by_so,
                                          po_log_prior=prior, now=now)
    ws = load_workbook(io.BytesIO(data))["SO-PO Log"]
    header = [c.value for c in ws[1]]
    return data, [dict(zip(header, r)) for r in ws.iter_rows(min_row=2, values_only=True)]


def _by_key(rows):
    return {(r["SO Number"], r["PO Number"] or ""): r for r in rows}


def test_new_so_without_po_gets_placeholder():
    _, rows = _run([_so("SO-1")], {}, None, T1)
    assert rows[0]["PO Status"] == "No PO yet"
    assert rows[0]["Last Change"] == "New SO - no PO yet"
    assert rows[0]["First Logged"] == T1


def test_placeholder_replaced_by_link_when_po_generated():
    data, _ = _run([_so("SO-1")], {}, None, T1)
    _, rows = _run([_so("SO-1")], {"SO-1": [_line()]}, data, T2)

    assert len(rows) == 1
    assert rows[0]["PO Number"] == "PO-001"
    assert rows[0]["Last Change"] == "PO linked (Draft) - was No PO yet"
    assert rows[0]["Last Changed"] == T2


def test_unchanged_link_keeps_its_last_changed_time():
    data, _ = _run([_so("SO-1")], {"SO-1": [_line()]}, None, T1)
    _, rows = _run([_so("SO-1")], {"SO-1": [_line()]}, data, T2)
    assert rows[0]["Last Changed"] == T1
    assert rows[0]["Last Change"] == "PO linked (Draft)"


def test_status_and_receipt_changes_are_stamped():
    data, _ = _run([_so("SO-1")], {"SO-1": [_line()]}, None, T1)
    _, rows = _run([_so("SO-1")], {"SO-1": [_line(status="Open", received=4)]}, data, T2)
    assert rows[0]["Last Changed"] == T2
    assert rows[0]["Last Change"] == "Status Draft -> Open; Received 0/1 -> 1/1"
    assert rows[0]["First Logged"] == T1


def test_history_survives_so_and_po_closing():
    data, _ = _run([_so("SO-1"), _so("SO-2")], {"SO-1": [_line()]}, None, T1)
    data, rows = _run([_so("SO-2")], {}, data, T2)  # SO-1 + its PO left BC's open set

    row = _by_key(rows)[("SO-1", "PO-001")]
    assert row["PO Status"] == "Closed in BC"
    assert row["Last Change"] == "SO and PO closed in BC"
    assert row["Customer"] == "Acme Doors"

    _, rows = _run([_so("SO-2")], {}, data, T3)  # stays put, not re-stamped
    assert _by_key(rows)[("SO-1", "PO-001")]["Last Changed"] == T2


def test_so_closed_without_ever_getting_a_po_is_recorded():
    data, _ = _run([_so("SO-1"), _so("SO-2")], {}, None, T1)
    _, rows = _run([_so("SO-2")], {}, data, T2)
    assert _by_key(rows)[("SO-1", "")]["PO Status"] == "SO closed - no PO"


def test_hand_typed_notes_carry_forward():
    data, _ = _run([_so("SO-1")], {"SO-1": [_line()]}, None, T1)
    wb = load_workbook(io.BytesIO(data))
    ws = wb["SO-PO Log"]
    notes_col = [c.value for c in ws[1]].index("Notes") + 1
    ws.cell(row=2, column=notes_col, value="Rush - customer called")
    buf = io.BytesIO()
    wb.save(buf)

    _, rows = _run([_so("SO-1")], {"SO-1": [_line(status="Open")]}, buf.getvalue(), T2)
    assert rows[0]["Notes"] == "Rush - customer called"


def test_po_referencing_unlogged_closed_so_is_ignored():
    """Old POs with typo'd / ancient SO refs shouldn't flood the log."""
    _, rows = _run([_so("SO-1")], {"SO-00173": [_line(po="PO-838")]}, None, T1)
    assert [r["SO Number"] for r in rows] == ["SO-1"]


def test_most_recently_changed_first():
    data, _ = _run([_so("SO-1"), _so("SO-2")], {"SO-1": [_line()], "SO-2": [_line(po="PO-002")]}, None, T1)
    _, rows = _run([_so("SO-1"), _so("SO-2")],
                   {"SO-1": [_line(status="Open")], "SO-2": [_line(po="PO-002")]}, data, T2)
    assert [r["SO Number"] for r in rows] == ["SO-1", "SO-2"]
