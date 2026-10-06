"""Upwardor order-acknowledgement intake.

Parser, mailbox poll (Graph and BC mocked), Production Schedule columns,
OData Vendor Order No. write, and the admin list. No live Graph, SharePoint,
or Business Central calls.
"""
from __future__ import annotations

import base64
import io
import shutil
import sys
from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.api.vendor_acks import get_current_admin, get_db, router as vendor_ack_router
from app.config import Settings, settings
from app.db.models import VendorOrderAck, VendorOrderAckSource
from app.integrations.bc.client import BusinessCentralClient
from app.services import scheduler_service as sched_mod
from app.services import vendor_ack_intake_service as intake_mod
from app.services.production_schedule_service import (
    PO_LINKS_HEADERS,
    production_schedule_service as schedule_svc,
)
from app.services.scheduler_service import SchedulerService, vendor_ack_interval_minutes
from app.services.vendor_ack_intake_service import (
    pdf_bytes_to_text,
    vendor_ack_intake_service as intake,
)
from app.services.vendor_ack_parser import (
    ack_cells_for_po,
    is_candidate_email,
    parse_acknowledgement,
    po_base,
    po_suffix,
)

FIXTURES = Path(__file__).resolve().parent / "fixtures" / "upwardor"
PDF_116307 = (FIXTURES / "s_ord_116307.txt").read_text()
PDF_116307_REV = (FIXTURES / "s_ord_116307_revised.txt").read_text()
PDF_116298 = (FIXTURES / "s_ord_116298_split.txt").read_text()
PDF_116299 = (FIXTURES / "s_ord_116299.txt").read_text()


def _db():
    # StaticPool keeps one connection so TestClient's worker thread sees the
    # same in-memory database the test seeded.
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    VendorOrderAck.__table__.create(engine)
    VendorOrderAckSource.__table__.create(engine)
    return sessionmaker(bind=engine)()


class FakeGraph:
    def __init__(self, emails, attachments):
        self.emails = emails  # mailbox -> [message]
        self.attachments = attachments  # (mailbox, message_id) -> [attachment]
        self.calls = []

    def get_recent_emails(self, mailbox, hours=24, max_count=50):
        self.calls.append(("list", mailbox, hours))
        return list(self.emails.get(mailbox, []))

    def get_message_attachments(self, mailbox, message_id):
        self.calls.append(("attachments", mailbox, message_id))
        return list(self.attachments.get((mailbox, message_id), []))


class FakeBC:
    def __init__(self):
        self.writes = []
        self.pos = {}
        self.fail = False

    def get_purchase_order_by_number(self, number):
        if self.fail:
            raise RuntimeError("bc down")
        return self.pos.get(number)

    def set_purchase_order_vendor_order_no(self, po_number, vendor_order_no):
        if self.fail:
            raise RuntimeError("bc down")
        self.writes.append((po_number, vendor_order_no))
        return {"number": po_number, "Vendor_Order_No": vendor_order_no}


def _email(message_id, subject, body, received, sender="andrew@upwardor.com"):
    return {
        "id": message_id,
        "subject": subject,
        "hasAttachments": True,
        "receivedDateTime": received,
        "body": {"contentType": "html", "content": body},
        "from": {"emailAddress": {"address": sender}},
    }


def _pdf_attachment(filename, key: bytes):
    return {
        "name": filename,
        "contentType": "application/pdf",
        "contentBytes": base64.b64encode(key).decode(),
    }


@pytest.fixture
def texts():
    return {}


@pytest.fixture
def fake_bc(monkeypatch):
    bc = FakeBC()
    bc.pos["PO-000956"] = {"id": "guid-956", "number": "PO-000956"}
    bc.pos["PO-000960"] = {"id": "guid-960", "number": "PO-000960"}
    bc.pos["PO-000961"] = {"id": "guid-961", "number": "PO-000961"}
    monkeypatch.setattr(intake_mod, "bc_client", bc)
    return bc


@pytest.fixture
def patched_settings(monkeypatch):
    monkeypatch.setattr(settings, "VENDOR_ACK_DRY_RUN", False)
    monkeypatch.setattr(settings, "VENDOR_ACK_BC_WRITEBACK", True)
    monkeypatch.setattr(
        settings, "VENDOR_ACK_INTAKE_MAILBOXES", "joey@opendc.ca,Finance@opendc.ca"
    )
    monkeypatch.setattr(settings, "VENDOR_ACK_INTAKE_LOOKBACK_HOURS", 72)
    monkeypatch.setattr(settings, "VENDOR_ACK_VENDOR_NO", "UPW")


def _patch_pdf(monkeypatch, texts):
    monkeypatch.setattr(intake_mod, "pdf_bytes_to_text", lambda data: texts.get(data, ""))


def _run(db, graph, texts, monkeypatch):
    monkeypatch.setattr(intake_mod, "graph_client", graph)
    _patch_pdf(monkeypatch, texts)
    return intake.process_new_acks(db)


# ── parser ──────────────────────────────────────────────────────────────

def test_andrew_confirmation_subject_and_pdf():
    parsed = parse_acknowledgement(
        subject="Confirmation Order#116307 / PO#000956",
        body="",
        filename="S-ORD116307.pdf",
        pdf_text=PDF_116307,
    )
    assert parsed.vendor_order_no == "S-ORD116307"
    assert parsed.our_po_number == "PO-000956"
    assert parsed.our_so_numbers == ["SO-001450"]
    assert parsed.status == "confirmed"
    assert parsed.completion_date == date(2026, 11, 20)
    assert parsed.po_source == "pdf"


def test_revised_subject_and_month_name_completion():
    parsed = parse_acknowledgement(
        subject="Revised Confirmation Order#116307 / PO#000956",
        filename="S-ORD116307.pdf",
        pdf_text=PDF_116307_REV,
    )
    assert parsed.status == "revised"
    assert parsed.completion_date == date(2026, 12, 1)
    assert parsed.our_po_number == "PO-000956"


def test_split_po_keeps_suffix_and_multiple_sos():
    parsed = parse_acknowledgement(
        subject="Confirmation Order#116298 / PO#000960(2)",
        filename="S-ORD116298.pdf",
        pdf_text=PDF_116298,
    )
    assert parsed.vendor_order_no == "S-ORD116298"
    assert parsed.our_po_number == "PO-000960(2)"
    assert po_suffix(parsed.our_po_number) == "2"
    assert po_base(parsed.our_po_number) == "PO-000960"
    assert parsed.our_so_numbers == ["SO-001455", "SO-001460"]


def test_cancelled_subject():
    parsed = parse_acknowledgement(
        subject="Cancelled Confirmation Order#116307 / PO#000956",
        filename="S-ORD116307.pdf",
        pdf_text=PDF_116307,
    )
    assert parsed.status == "cancelled"
    assert parsed.our_po_number == "PO-000956"


def test_subject_only_when_pdf_text_missing():
    parsed = parse_acknowledgement(
        subject="Confirmation Order#116307 / PO#000956",
        filename="S-ORD116307.pdf",
        pdf_text="",
    )
    assert parsed.vendor_order_no == "S-ORD116307"
    assert parsed.our_po_number == "PO-000956"
    assert parsed.completion_date is None
    assert parsed.our_so_numbers == []


def test_no_po_is_pending():
    parsed = parse_acknowledgement(
        subject="FW: docs",
        body="attached order confirmation",
        filename="S-ORD116400.pdf",
        pdf_text="S-ORD116400\nCompletion Date: 11/01/2026\n",
    )
    assert parsed.status == "pending"
    assert parsed.document_status == "confirmed"
    assert parsed.our_po_number is None
    assert parsed.completion_date == date(2026, 11, 1)


def test_unrelated_pdf_is_not_an_ack():
    assert parse_acknowledgement(
        subject="FW: docs", body="attached order confirmation",
        filename="lunch.pdf", pdf_text="Lunch menu for Friday",
    ) is None


def test_manveer_candidate_and_andrew_candidate():
    assert is_candidate_email(
        "FW: from Manveer", "<p>Please see the attached order confirmation</p>"
    )
    assert is_candidate_email("Confirmation Order#116307 / PO#000956", "")
    assert not is_candidate_email("Invoice 4401", "Please pay the attached invoice")


def test_ack_cells_label_split_without_replacing_primary():
    cells = ack_cells_for_po("PO-000960", [
        {
            "vendor_order_no": "S-ORD116299",
            "our_po_number": "PO-000960",
            "status": "confirmed",
            "completion_date": date(2026, 12, 1),
            # 05:30 UTC is still the previous evening in Edmonton.
            "received_at": datetime(2026, 10, 7, 5, 30),
        },
        {
            "vendor_order_no": "S-ORD116298",
            "our_po_number": "PO-000960(2)",
            "status": "revised",
            "completion_date": date(2026, 12, 5),
            "received_at": datetime(2026, 10, 3, 15, 0),
        },
    ])
    assert cells["vendor_ack"] == "S-ORD116299; S-ORD116298 (2)"
    assert cells["ack_status"] == "Confirmed; Revised (2)"
    assert cells["completion_date"] == date(2026, 12, 1)
    assert cells["ack_received"] == date(2026, 10, 6)


def test_default_mailboxes_and_interval():
    assert Settings.model_fields["VENDOR_ACK_INTAKE_MAILBOXES"].default == (
        "joey@opendc.ca,Finance@opendc.ca"
    )
    assert Settings.model_fields["VENDOR_ACK_INTAKE_ENABLED"].default is False
    assert vendor_ack_interval_minutes(20) == 20
    assert vendor_ack_interval_minutes(5) == 15
    assert vendor_ack_interval_minutes(90) == 30


# ── intake ──────────────────────────────────────────────────────────────

def test_andrew_email_writes_vendor_order_no(monkeypatch, fake_bc, patched_settings, texts):
    texts[b"pdf-116307"] = PDF_116307
    graph = FakeGraph(
        {"joey@opendc.ca": [_email(
            "msg-1",
            "Confirmation Order#116307 / PO#000956",
            "<p>Confirmation attached</p>",
            "2026-10-01T15:04:05Z",
        )]},
        {("joey@opendc.ca", "msg-1"): [_pdf_attachment("S-ORD116307.pdf", b"pdf-116307")]},
    )
    db = _db()
    summary = _run(db, graph, texts, monkeypatch)
    assert summary["parsed"] == 1
    assert summary["bc_written"] == 1
    assert fake_bc.writes == [("PO-000956", "S-ORD116307")]
    row = db.query(VendorOrderAck).one()
    assert row.vendor_no == "UPW"
    assert row.vendor_order_no == "S-ORD116307"
    assert row.our_po_number == "PO-000956"
    assert row.our_so_numbers == ["SO-001450"]
    assert row.status == "confirmed"
    assert row.completion_date == date(2026, 11, 20)
    assert row.bc_write_status == "written"
    assert row.bc_po_id == "guid-956"
    assert row.mailbox == "joey@opendc.ca"

    again = _run(db, graph, texts, monkeypatch)
    assert again["duplicate"] == 1
    assert again["parsed"] == 0
    assert db.query(VendorOrderAck).count() == 1
    assert fake_bc.writes == [("PO-000956", "S-ORD116307")]


def test_revised_email_upserts_latest_wins_even_if_newer_is_listed_first(
    monkeypatch, fake_bc, patched_settings, texts,
):
    texts[b"pdf-old"] = PDF_116307
    texts[b"pdf-new"] = PDF_116307_REV
    graph = FakeGraph(
        {"joey@opendc.ca": [
            _email(
                "msg-new", "Revised Confirmation Order#116307 / PO#000956",
                "revised", "2026-10-03T18:00:00Z",
            ),
            _email(
                "msg-old", "Confirmation Order#116307 / PO#000956",
                "original", "2026-10-01T15:00:00Z",
            ),
        ]},
        {
            ("joey@opendc.ca", "msg-new"): [_pdf_attachment("S-ORD116307.pdf", b"pdf-new")],
            ("joey@opendc.ca", "msg-old"): [_pdf_attachment("S-ORD116307.pdf", b"pdf-old")],
        },
    )
    db = _db()
    summary = _run(db, graph, texts, monkeypatch)
    assert summary["parsed"] == 2
    assert summary["updated"] == 1
    row = db.query(VendorOrderAck).one()
    assert row.status == "revised"
    assert row.completion_date == date(2026, 12, 1)
    assert row.source_email_id == "msg-new"
    assert db.query(VendorOrderAckSource).count() == 2
    assert [call[1] for call in fake_bc.writes] == ["S-ORD116307", "S-ORD116307"]


def test_split_does_not_overwrite_primary_po(monkeypatch, fake_bc, patched_settings, texts):
    texts[b"primary"] = PDF_116307.replace("000956", "000960").replace("S-ORD116307", "S-ORD116310")
    texts[b"split"] = PDF_116298
    graph = FakeGraph(
        {"Finance@opendc.ca": [
            _email("m1", "Confirmation Order#116310 / PO#000960", "", "2026-10-01T12:00:00Z"),
            _email("m2", "Confirmation Order#116298 / PO#000960(2)", "", "2026-10-01T13:00:00Z"),
        ]},
        {
            ("Finance@opendc.ca", "m1"): [_pdf_attachment("S-ORD116310.pdf", b"primary")],
            ("Finance@opendc.ca", "m2"): [_pdf_attachment("S-ORD116298.pdf", b"split")],
        },
    )
    db = _db()
    _run(db, graph, texts, monkeypatch)
    primary = db.query(VendorOrderAck).filter_by(vendor_order_no="S-ORD116310").one()
    split = db.query(VendorOrderAck).filter_by(vendor_order_no="S-ORD116298").one()
    assert primary.our_po_number == "PO-000960"
    assert primary.status == "confirmed"
    assert primary.bc_write_status == "written"
    assert split.our_po_number == "PO-000960(2)"
    assert split.bc_write_status == "skipped"
    assert "primary PO" in (split.bc_write_error or "")
    assert fake_bc.writes == [("PO-000960", "S-ORD116310")]


def test_manveer_multi_pdf(monkeypatch, fake_bc, patched_settings, texts):
    texts[b"a"] = PDF_116298
    texts[b"b"] = PDF_116299
    graph = FakeGraph(
        {"joey@opendc.ca": [_email(
            "mv-1",
            "FW: see below",
            "<div>attached order confirmation</div>",
            "2026-10-04T16:00:00Z",
            sender="manveer@opendc.ca",
        )]},
        {("joey@opendc.ca", "mv-1"): [
            _pdf_attachment("Acknowledgement_S-ORD116298.pdf", b"a"),
            _pdf_attachment("S-ORD116299.pdf", b"b"),
            _pdf_attachment("photo.jpg", b"nope"),
        ]},
    )
    db = _db()
    summary = _run(db, graph, texts, monkeypatch)
    assert summary["parsed"] == 2
    numbers = {row.vendor_order_no for row in db.query(VendorOrderAck).all()}
    assert numbers == {"S-ORD116298", "S-ORD116299"}
    written = db.query(VendorOrderAck).filter_by(vendor_order_no="S-ORD116299").one()
    assert written.our_po_number == "PO-000961"
    assert written.our_so_numbers == ["SO-001470"]
    assert written.completion_date == date(2026, 12, 15)
    assert written.bc_write_status == "written"
    assert ("PO-000961", "S-ORD116299") in fake_bc.writes
    assert all(po != "PO-000960" for po, _ack in fake_bc.writes)


def test_pending_when_po_missing_and_non_ack_pdf_is_idempotent(
    monkeypatch, fake_bc, patched_settings, texts,
):
    texts[b"pending"] = "S-ORD116400\nCompletion Date: 11/01/2026\n"
    texts[b"menu"] = "Lunch menu"
    graph = FakeGraph(
        {"joey@opendc.ca": [_email(
            "mv-2", "hello", "attached order confirmation", "2026-10-05T12:00:00Z",
        )]},
        {("joey@opendc.ca", "mv-2"): [
            _pdf_attachment("S-ORD116400.pdf", b"pending"),
            _pdf_attachment("notes.pdf", b"menu"),
        ]},
    )
    db = _db()
    first = _run(db, graph, texts, monkeypatch)
    assert first["pending"] == 1
    assert first["skipped"] == 1
    assert first["pending_orders"] == ["S-ORD116400"]
    row = db.query(VendorOrderAck).one()
    assert row.status == "pending"
    assert row.bc_write_status == "skipped"
    assert fake_bc.writes == []
    second = _run(db, graph, texts, monkeypatch)
    assert second["duplicate"] == 2
    assert second["parsed"] == 0
    assert db.query(VendorOrderAck).count() == 1


def test_bc_failure_does_not_drop_the_ack_and_retries(monkeypatch, fake_bc, patched_settings, texts):
    texts[b"pdf-116307"] = PDF_116307
    graph = FakeGraph(
        {"joey@opendc.ca": [_email(
            "msg-1", "Confirmation Order#116307 / PO#000956", "", "2026-10-01T15:00:00Z",
        )]},
        {("joey@opendc.ca", "msg-1"): [_pdf_attachment("S-ORD116307.pdf", b"pdf-116307")]},
    )
    fake_bc.fail = True
    db = _db()
    summary = _run(db, graph, texts, monkeypatch)
    assert summary["parsed"] == 1
    assert summary["bc_failed"] == 1
    row = db.query(VendorOrderAck).one()
    assert row.vendor_order_no == "S-ORD116307"
    assert row.bc_write_status == "failed"
    assert "bc down" in row.bc_write_error

    fake_bc.fail = False
    retried = _run(db, graph, texts, monkeypatch)
    assert retried["duplicate"] == 1
    assert retried["bc_retried"] == 1
    db.refresh(row)
    assert row.bc_write_status == "written"
    assert fake_bc.writes == [("PO-000956", "S-ORD116307")]


def test_older_failure_does_not_clobber_a_later_ack(monkeypatch, fake_bc, patched_settings, texts):
    texts[b"old"] = PDF_116307
    texts[b"new"] = PDF_116307.replace("S-ORD116307", "S-ORD116500")
    graph = FakeGraph(
        {"joey@opendc.ca": [
            _email("old", "Confirmation Order#116307 / PO#000956", "", "2026-10-01T12:00:00Z"),
            _email("new", "Confirmation Order#116500 / PO#000956", "", "2026-10-02T12:00:00Z"),
        ]},
        {
            ("joey@opendc.ca", "old"): [_pdf_attachment("S-ORD116307.pdf", b"old")],
            ("joey@opendc.ca", "new"): [_pdf_attachment("S-ORD116500.pdf", b"new")],
        },
    )
    db = _db()
    _run(db, graph, texts, monkeypatch)
    assert fake_bc.writes[-1] == ("PO-000956", "S-ORD116500")
    older = db.query(VendorOrderAck).filter_by(vendor_order_no="S-ORD116307").one()
    older.bc_write_status = "failed"
    db.commit()
    intake._retry_failed_writes(db, dry_run=False)
    db.refresh(older)
    assert older.bc_write_status == "skipped"
    assert fake_bc.writes[-1] == ("PO-000956", "S-ORD116500")


def test_dry_run_skips_bc(monkeypatch, fake_bc, patched_settings, texts):
    texts[b"pdf-116307"] = PDF_116307
    graph = FakeGraph(
        {"joey@opendc.ca": [_email(
            "msg-1", "Confirmation Order#116307 / PO#000956", "", "2026-10-01T15:00:00Z",
        )]},
        {("joey@opendc.ca", "msg-1"): [_pdf_attachment("S-ORD116307.pdf", b"pdf-116307")]},
    )
    db = _db()
    monkeypatch.setattr(intake_mod, "graph_client", graph)
    _patch_pdf(monkeypatch, texts)
    summary = intake.process_new_acks(db, dry_run=True)
    assert summary["dry_run"] is True
    assert summary["parsed"] == 1
    row = db.query(VendorOrderAck).one()
    assert row.bc_write_status == "dry_run"
    assert fake_bc.writes == []


def test_both_mailboxes_are_polled(monkeypatch, fake_bc, patched_settings, texts):
    texts[b"pdf-116307"] = PDF_116307
    graph = FakeGraph(
        {
            "joey@opendc.ca": [],
            "Finance@opendc.ca": [_email(
                "fin-1", "Confirmation Order#116307 / PO#000956", "", "2026-10-02T15:00:00Z",
            )],
        },
        {("Finance@opendc.ca", "fin-1"): [_pdf_attachment("S-ORD116307.pdf", b"pdf-116307")]},
    )
    db = _db()
    summary = _run(db, graph, texts, monkeypatch)
    assert summary["mailboxes"] == ["joey@opendc.ca", "Finance@opendc.ca"]
    assert {call[1] for call in graph.calls if call[0] == "list"} == {
        "joey@opendc.ca", "Finance@opendc.ca",
    }
    assert db.query(VendorOrderAck).one().mailbox == "Finance@opendc.ca"


def test_non_candidate_mail_is_not_opened(monkeypatch, fake_bc, patched_settings, texts):
    graph = FakeGraph(
        {"joey@opendc.ca": [_email(
            "inv-1", "Invoice 4401", "Please pay the attached invoice", "2026-10-01T12:00:00Z",
        )]},
        {("joey@opendc.ca", "inv-1"): [_pdf_attachment("invoice.pdf", b"inv")]},
    )
    db = _db()
    summary = _run(db, graph, texts, monkeypatch)
    assert summary["processed"] == 0
    assert db.query(VendorOrderAck).count() == 0
    assert not any(call[0] == "attachments" for call in graph.calls)


# ── production schedule sheet ───────────────────────────────────────────

def _so(number, items, customer="Acme Doors"):
    return {
        "number": number, "customerName": customer, "externalDocumentNumber": "TAG",
        "orderDate": "2026-09-30",
        "salesOrderLines": [
            {"lineType": "Item", "lineObjectNumber": item, "quantity": qty, "description": item}
            for item, qty in items
        ],
    }


def _po_line(item, qty, po, received=0):
    return {
        "lineType": "Item", "lineObjectNumber": item, "quantity": qty,
        "receivedQuantity": received, "expectedReceiptDate": "2026-10-15",
        "_po_number": po, "_po_status": "Open", "_vendor": "UPWARDOR",
    }


def test_purchase_orders_sheet_shows_ack_columns():
    orders = [_so("SO-001450", [("PN45-24400-1400", 4)])]
    po_lines = {"SO-001450": [
        _po_line("PN45-24400-1400", 4, "PO-000956"),
        _po_line("PN45-24400-1400", 2, "PO-000960"),
    ]}
    facts = schedule_svc.compute_so_facts(orders, po_lines)
    acks = [
        {
            "vendor_order_no": "S-ORD116307", "our_po_number": "PO-000956",
            "status": "confirmed", "completion_date": date(2026, 11, 20),
            "received_at": datetime(2026, 10, 1, 15, 4),
        },
        {
            "vendor_order_no": "S-ORD116299", "our_po_number": "PO-000960",
            "status": "confirmed", "completion_date": date(2026, 12, 1),
            "received_at": datetime(2026, 10, 2, 15, 0),
        },
        {
            "vendor_order_no": "S-ORD116298", "our_po_number": "PO-000960(2)",
            "status": "revised", "completion_date": date(2026, 12, 5),
            "received_at": datetime(2026, 10, 3, 15, 0),
        },
    ]
    data, _, _ = schedule_svc.build_workbook_bytes(
        orders, {}, facts, po_lines_by_so=po_lines, vendor_acks=acks,
    )
    ws = load_workbook(io.BytesIO(data))["Purchase Orders"]
    assert [c.value for c in ws[1]] == PO_LINKS_HEADERS
    rows = {row[2]: row for row in ws.iter_rows(min_row=2, values_only=True)}
    assert rows["PO-000956"][8] == "S-ORD116307"
    assert rows["PO-000956"][9] == "Confirmed"
    completion = rows["PO-000956"][11]
    assert (completion.date() if isinstance(completion, datetime) else completion) == date(2026, 11, 20)
    assert rows["PO-000960"][8] == "S-ORD116299; S-ORD116298 (2)"
    assert "Revised (2)" in rows["PO-000960"][9]
    # Received column stays where the existing sheet left it.
    assert rows["PO-000956"][6] == "0/1"


def test_schedule_ack_lookup_failure_is_blank(monkeypatch):
    def _boom():
        raise RuntimeError("db down")

    monkeypatch.setattr("app.db.database.SessionLocal", _boom)
    assert schedule_svc._load_vendor_acks() == []


# ── BC OData write ──────────────────────────────────────────────────────

def test_vendor_order_no_patches_odata_purchase_order(monkeypatch):
    bc = BusinessCentralClient.__new__(BusinessCentralClient)
    bc.odata_url = "https://bc.example/ODataV4"
    bc._get_access_token = lambda: "tok"
    bc._odata_v4_company_segment = lambda company_id=None: "Company('OPENDC')"
    monkeypatch.setattr(settings, "VENDOR_ACK_BC_ODATA_ENTITY", "PurchaseOrder")
    monkeypatch.setattr(settings, "VENDOR_ACK_BC_ODATA_FIELD", "Vendor_Order_No")
    seen = {}

    class Resp:
        def __init__(self, status, payload):
            self.status_code = status
            self._payload = payload
            self.text = ""
            self.content = b"{}"

        def json(self):
            return self._payload

    def fake_get(url, headers=None, timeout=None):
        seen["get"] = url
        return Resp(200, {"value": [
            {"Document_Type": "Quote", "No": "PO-000956"},
            {
                "Document_Type": "Order",
                "No": "PO-000956",
                "@odata.etag": 'W/"abc"',
                "@odata.editLink": "https://bc.example/ODataV4/Company('OPENDC')/PurchaseOrder(Document_Type='Order',No='PO-000956')",
            },
        ]})

    def fake_patch(url, json=None, headers=None, timeout=None):
        seen["patch_url"] = url
        seen["patch_json"] = json
        seen["if_match"] = headers.get("If-Match")
        return Resp(200, {"No": "PO-000956", "Vendor_Order_No": json["Vendor_Order_No"]})

    import app.integrations.bc.client as bc_mod
    monkeypatch.setattr(bc_mod.requests, "get", fake_get)
    monkeypatch.setattr(bc_mod.requests, "patch", fake_patch)

    result = bc.set_purchase_order_vendor_order_no("PO-000956", "S-ORD116307")
    assert "No eq 'PO-000956'" in seen["get"]
    assert seen["patch_json"] == {"Vendor_Order_No": "S-ORD116307"}
    assert seen["if_match"] == 'W/"abc"'
    assert "Document_Type='Order'" in seen["patch_url"]
    assert result["Vendor_Order_No"] == "S-ORD116307"


def test_missing_odata_row_raises(monkeypatch):
    bc = BusinessCentralClient.__new__(BusinessCentralClient)
    bc.odata_url = "https://bc.example/ODataV4"
    bc._get_access_token = lambda: "tok"
    bc._odata_v4_company_segment = lambda company_id=None: "Company('OPENDC')"

    class Resp:
        status_code = 200
        text = ""
        content = b"{}"

        def json(self):
            return {"value": []}

    import app.integrations.bc.client as bc_mod
    monkeypatch.setattr(bc_mod.requests, "get", lambda *a, **k: Resp())
    with pytest.raises(LookupError):
        bc.set_purchase_order_vendor_order_no("PO-000956", "S-ORD116307")


# ── scheduler + admin API ───────────────────────────────────────────────

def test_scheduler_job_skips_when_disabled(monkeypatch):
    monkeypatch.setattr(settings, "VENDOR_ACK_INTAKE_ENABLED", False)
    called = {}

    def _nope(*args, **kwargs):
        called["ran"] = True
        return {}

    monkeypatch.setattr(intake_mod.vendor_ack_intake_service, "process_new_acks", _nope)
    SchedulerService()._vendor_ack_intake_job()
    assert "ran" not in called


def test_scheduler_job_runs_when_enabled(monkeypatch):
    monkeypatch.setattr(settings, "VENDOR_ACK_INTAKE_ENABLED", True)
    db = _db()
    monkeypatch.setattr(sched_mod, "SessionLocal", lambda: db)
    seen = {}

    def _run_acks(session, dry_run=None):
        seen["db"] = session
        seen["dry_run"] = dry_run
        return {"parsed": 0}

    monkeypatch.setattr(intake_mod.vendor_ack_intake_service, "process_new_acks", _run_acks)
    SchedulerService()._vendor_ack_intake_job()
    assert seen["db"] is db
    assert seen["dry_run"] is None


def test_admin_list_and_dry_run(monkeypatch):
    db = _db()
    db.add(VendorOrderAck(
        source_email_id="e1",
        attachment_filename="S-ORD116400.pdf",
        vendor_no="UPW",
        vendor_order_no="S-ORD116400",
        status="pending",
        our_so_numbers=[],
        bc_write_status="skipped",
    ))
    db.commit()

    app = FastAPI()
    app.include_router(vendor_ack_router)

    def _override_db():
        yield db

    app.dependency_overrides[get_db] = _override_db
    app.dependency_overrides[get_current_admin] = lambda: SimpleNamespace(id=1, is_active=True)
    client = TestClient(app)

    bare = FastAPI()
    bare.include_router(vendor_ack_router)
    assert TestClient(bare).get("/api/admin/vendor-acks").status_code in (401, 403)

    listed = client.get("/api/admin/vendor-acks", params={"status": "pending"})
    assert listed.status_code == 200
    body = listed.json()
    assert body["counts"]["pending"] == 1
    assert body["acknowledgements"][0]["vendor_order_no"] == "S-ORD116400"

    seen = {}

    def _fake_run(session, dry_run=None):
        seen["dry_run"] = dry_run
        return {"parsed": 0, "dry_run": dry_run}

    monkeypatch.setattr(intake_mod.vendor_ack_intake_service, "process_new_acks", _fake_run)
    ran = client.post("/api/admin/vendor-acks/run", params={"dry_run": "true"})
    assert ran.status_code == 200
    assert seen["dry_run"] is True
    assert ran.json()["dry_run"] is True


def test_pdftotext_round_trip_on_generated_pdf():
    if not shutil.which("pdftotext"):
        pytest.skip("pdftotext is not installed")
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.multi_cell(0, 6, PDF_116307)
    raw = pdf.output()
    extracted = pdf_bytes_to_text(bytes(raw))
    parsed = parse_acknowledgement(
        subject="Confirmation Order#116307 / PO#000956",
        filename="S-ORD116307.pdf",
        pdf_text=extracted,
    )
    assert parsed is not None
    assert parsed.vendor_order_no == "S-ORD116307"
    assert parsed.our_po_number == "PO-000956"
    assert parsed.completion_date == date(2026, 11, 20)
    assert parsed.our_so_numbers == ["SO-001450"]
