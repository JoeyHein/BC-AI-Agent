"""Partial Build on the production schedule (Joey, 2026-10-08).

Aluminum sections are built at OpenDC; the rest of the order is bought
complete from Upwardor. Fulfillment gains "Partial Build", seeded only
when the cell is still the untouched default Buy Complete. In-house Build
tracks that half. Read-back stays by header name, and the SO-PO Log is
left intact when the Schedule column set grows.
"""
import io
import os
import sys
from datetime import date, datetime

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from openpyxl import Workbook, load_workbook

from app.services.production_schedule_service import (
    INHOUSE_BUILD_HEADER,
    SCHEDULE_HEADERS,
    production_schedule_service as svc,
    is_inhouse_aluminum_line,
    _SORecord,
)


def _so(number, items, customer="Acme Doors"):
    lines = []
    for spec in items:
        item, qty = spec[0], spec[1]
        description = spec[2] if len(spec) > 2 else item
        category = spec[3] if len(spec) > 3 else ""
        lines.append({
            "lineType": "Item",
            "lineObjectNumber": item,
            "quantity": qty,
            "description": description,
            "itemCategoryCode": category,
        })
    return {
        "number": number, "customerName": customer, "externalDocumentNumber": "TAG",
        "orderDate": "2026-09-30", "salesOrderLines": lines,
    }


def _po_line(item, qty, received=0, po="PO-001", status="Open", expected="2026-10-15"):
    return {
        "lineType": "Item", "lineObjectNumber": item, "quantity": qty,
        "receivedQuantity": received, "expectedReceiptDate": expected,
        "_po_number": po, "_po_status": status, "_vendor": "UPWARDOR",
    }


def _schedule(orders, records=None, po_lines_by_so=None, po_log_prior=None, now=None):
    facts = svc.compute_so_facts(orders, po_lines_by_so or {})
    data, _, _ = svc.build_workbook_bytes(
        orders, records or {}, facts, po_lines_by_so=po_lines_by_so,
        po_log_prior=po_log_prior, now=now,
    )
    wb = load_workbook(io.BytesIO(data))
    ws = wb["Schedule"]
    header = [c.value for c in ws[1]]
    rows = [dict(zip(header, row)) for row in ws.iter_rows(min_row=2, values_only=True)]
    return rows, data, wb


# ── detection rule ────────────────────────────────────────────────────

def test_section_prefixes_match_and_lookalikes_do_not():
    assert is_inhouse_aluminum_line("PN80-21100-1002")
    assert is_inhouse_aluminum_line("pn10-21200010-0802")
    assert is_inhouse_aluminum_line("PN12-21200310-0802")
    assert is_inhouse_aluminum_line("PN97-21200010-0802")
    assert is_inhouse_aluminum_line("PN20-2100010-0802")
    assert is_inhouse_aluminum_line("PN70-21100010-0802")
    assert is_inhouse_aluminum_line("PN80")

    assert not is_inhouse_aluminum_line("PN45-24400-1400")
    assert not is_inhouse_aluminum_line("PN46-24400-1400")
    assert not is_inhouse_aluminum_line("PN100-1")
    assert not is_inhouse_aluminum_line("HK02-14120-RC")
    assert not is_inhouse_aluminum_line("TR02-1")
    assert not is_inhouse_aluminum_line("SP12-1")
    assert not is_inhouse_aluminum_line("GK17-25100-00", "GLAZING KIT, ALUM, SINGLE (3MM), CLEAR")
    assert not is_inhouse_aluminum_line("OP20-00001")
    assert not is_inhouse_aluminum_line("AL10-00002-00", "ALUMINIUM COIL, .032 X 24\"")
    assert not is_inhouse_aluminum_line("AL91-67400-00", 'CENTER STILE 2" AL976 AS-57674 MILL')
    assert not is_inhouse_aluminum_line("PL10-1", 'ASTRAGAL, 3" ALUMINIUM DOOR BOTTOM RUBBER')


def test_description_and_category_backstops():
    assert is_inhouse_aluminum_line("CUSTOM-1", "SECTION, PANORAMA, 10' CLEAR")
    assert is_inhouse_aluminum_line("CUSTOM-2", "Full view V130G section")
    assert is_inhouse_aluminum_line("CUSTOM-3", "SECTION, AL976, NO GLASS, CLEAR ANO")
    assert is_inhouse_aluminum_line("CUSTOM-4", "SOLALITE frame")
    assert is_inhouse_aluminum_line("CUSTOM-5", "", category="PANORAMA")
    assert is_inhouse_aluminum_line("CUSTOM-6", "", category="AL97621CA")
    assert is_inhouse_aluminum_line("CUSTOM-7", "", category="ALV1324WH")
    assert not is_inhouse_aluminum_line("CUSTOM-8", "SECTION, TX 450, WHITE")
    assert not is_inhouse_aluminum_line("CUSTOM-9", "", category="COMPONENT")
    assert not is_inhouse_aluminum_line("CUSTOM-10", "", category="ALGK")


def test_compute_so_facts_collects_aluminum_lines_only():
    facts = svc.compute_so_facts([_so("SO-1", [
        ("PN80-21100-1002", 4, "SECTION, PANORAMA"),
        ("PN45-24400-1400", 8, "SECTION, TX 450"),
        ("HK02-1", 1, "HARDWARE KIT"),
        ("GK17-25100-00", 4, "GLAZING KIT, ALUM"),
    ])], {})
    items = [ln["item"] for ln in facts["SO-1"]["aluminum_lines"]]
    assert items == ["PN80-21100-1002"]
    assert facts["SO-1"]["window_kit_lines"][0]["item"] == "GK17-25100-00"


# ── seeding ───────────────────────────────────────────────────────────

def test_new_aluminum_so_is_partial_build_and_po_status_still_tracks_the_buy():
    orders = [_so("SO-1328", [
        ("PN80-21100-1002", 4, "SECTION, PANORAMA"),
        ("PN45-24400-1400", 8, "polycarb panel"),
        ("HK02-1", 1, "HARDWARE KIT"),
        ("TR02-1", 2, "TRACK"),
    ], customer="GNB Doors")]
    po = {"SO-1328": [
        _po_line("PN45-24400-1400", 8),
        _po_line("HK02-1", 1),
        _po_line("TR02-1", 2),
    ]}
    rows, _, _ = _schedule(orders, po_lines_by_so=po)
    assert rows[0]["Fulfillment"] == "Partial Build"
    assert rows[0][INHOUSE_BUILD_HEADER] == "Not Started"
    assert rows[0]["Emergency Build"] is None
    assert rows[0]["Order Status"] == "Ordered"
    assert rows[0]["Upwardor PO #"] == "PO-001"
    assert rows[0]["Customer Name"] == "GNB Doors"


def test_so_without_aluminum_stays_buy_complete():
    rows, _, _ = _schedule([_so("SO-1", [("PN45-24400-1400", 4), ("HK02-1", 1)])])
    assert rows[0]["Fulfillment"] == "Buy Complete"
    assert rows[0][INHOUSE_BUILD_HEADER] is None


def test_untouched_buy_complete_on_a_v4_row_is_upgraded():
    """A record read from a sheet that predates Partial Build has
    fulfillment_explicit False, so Buy Complete is still the default."""
    prior = _SORecord("SO-1299")
    prior.fulfillment = "Buy Complete"
    prior.shipping_status = "Ready to Ship"
    prior.po_date = date(2026, 9, 2)
    rows, _, _ = _schedule(
        [_so("SO-1299", [("PN80-24100-1202", 4, "SECTION, PANORAMA")], customer="Panorama job")],
        records={"SO-1299": prior},
    )
    assert rows[0]["Fulfillment"] == "Partial Build"
    assert rows[0][INHOUSE_BUILD_HEADER] == "Not Started"
    assert rows[0]["Shipping Status"] == "Ready to Ship"
    assert rows[0]["PO Date"].date() == date(2026, 9, 2)


def test_hand_set_emergency_build_is_not_overwritten():
    prior = _SORecord("SO-1308")
    prior.fulfillment = "Emergency Build"
    prior.fulfillment_explicit = True
    prior.emergency_build = "In Production"
    rows, _, _ = _schedule(
        [_so("SO-1308", [("PN80-21100-0802", 2, "SECTION, PANORAMA")], customer="Silver Spray")],
        records={"SO-1308": prior},
    )
    assert rows[0]["Fulfillment"] == "Emergency Build"
    assert rows[0]["Emergency Build"] == "In Production"
    assert rows[0][INHOUSE_BUILD_HEADER] is None


def test_hand_set_buy_complete_is_not_overwritten_once_partial_build_exists():
    prior = _SORecord("SO-1")
    prior.fulfillment = "Buy Complete"
    prior.fulfillment_explicit = True
    rows, _, _ = _schedule(
        [_so("SO-1", [("PN97-21200010-0802", 2, "SECTION, AL976")])],
        records={"SO-1": prior},
    )
    assert rows[0]["Fulfillment"] == "Buy Complete"
    assert rows[0][INHOUSE_BUILD_HEADER] is None


def test_hand_set_inhouse_status_carries_forward():
    prior = _SORecord("SO-1")
    prior.fulfillment = "Partial Build"
    prior.fulfillment_explicit = True
    prior.inhouse_build = "In Production"
    rows, _, _ = _schedule(
        [_so("SO-1", [("PN10-21200010-0802", 2, "SECTION, V130G")])],
        records={"SO-1": prior},
    )
    assert rows[0]["Fulfillment"] == "Partial Build"
    assert rows[0][INHOUSE_BUILD_HEADER] == "In Production"


def test_inhouse_build_clears_when_flipped_off_partial():
    prior = _SORecord("SO-1")
    prior.fulfillment = "Buy Complete"
    prior.fulfillment_explicit = True
    prior.inhouse_build = "Complete"
    rows, _, _ = _schedule(
        [_so("SO-1", [("PN80-21100-0802", 1)])],
        records={"SO-1": prior},
    )
    assert rows[0][INHOUSE_BUILD_HEADER] is None


# ── read-back and v4 migration ────────────────────────────────────────

def test_new_column_round_trips_by_header_name():
    prior = _SORecord("SO-1")
    prior.fulfillment = "Partial Build"
    prior.fulfillment_explicit = True
    prior.inhouse_build = "Complete"
    prior.shipping_status = "Ready to Ship"
    prior.po_date = date(2026, 9, 1)
    _, data, _ = _schedule(
        [_so("SO-1", [("PN80-21100-0802", 2, "SECTION, PANORAMA")])],
        records={"SO-1": prior},
    )

    rec = svc.parse_records_from_bytes(data)["SO-1"]
    assert rec.fulfillment == "Partial Build"
    assert rec.fulfillment_explicit is True
    assert rec.inhouse_build == "Complete"
    assert rec.shipping_status == "Ready to Ship"
    assert rec.po_date == date(2026, 9, 1)

    # Second refresh must not put Not Started back over Complete.
    rows, _, _ = _schedule(
        [_so("SO-1", [("PN80-21100-0802", 2, "SECTION, PANORAMA")])],
        records={"SO-1": rec},
    )
    assert rows[0][INHOUSE_BUILD_HEADER] == "Complete"
    assert rows[0]["Fulfillment"] == "Partial Build"


def test_shuffled_headers_still_read_inhouse_build():
    wb = Workbook()
    ws = wb.active
    ws.title = "Schedule"
    ws.append(["Shipping Status", INHOUSE_BUILD_HEADER, "SO Number", "Fulfillment"])
    ws.append(["Shipped", "In Production", "SO-1", "Partial Build"])
    buf = io.BytesIO()
    wb.save(buf)

    rec = svc.parse_records_from_bytes(buf.getvalue())["SO-1"]
    assert rec.fulfillment == "Partial Build"
    assert rec.fulfillment_explicit is True
    assert rec.inhouse_build == "In Production"
    assert rec.shipping_status == "Shipped"


def test_v4_workbook_upgrades_default_and_keeps_hand_set_and_po_log():
    """The live file before this column: no In-house Build header. Buy
    Complete on an aluminum SO is the untouched default. Emergency Build
    stays. SO-PO Log notes survive the rewrite."""
    old_headers = [h for h in SCHEDULE_HEADERS if h != INHOUSE_BUILD_HEADER]
    wb = Workbook()
    ws = wb.active
    ws.title = "Schedule"
    ws.append(old_headers)

    def put(so, customer, fulfillment, shipping, order_status):
        values = {h: None for h in old_headers}
        values.update({
            "SO Number": so,
            "Customer Name": customer,
            "Fulfillment": fulfillment,
            "Shipping Status": shipping,
            "Order Status": order_status,
            "PO Date": date(2026, 9, 15),
        })
        ws.append([values[h] for h in old_headers])

    put("SO-1328", "GNB Doors", "Buy Complete", "Not Ready", "Shipped by Vendor")
    put("SO-1308", "Silver Spray", "Emergency Build", "Ready to Ship", "Ordered")
    # Pull from stock is a hand-set on the same row as the emergency flip.
    pull_col = old_headers.index("Pull from stock")
    # Row 1 is the header, row 2 is SO-1328, row 3 is SO-1308.
    ws.cell(row=3, column=pull_col + 1, value="Yes")

    log = wb.create_sheet("SO-PO Log")
    log.append(["SO Number", "Customer", "PO Number", "Vendor", "PO Status", "PO Date",
                "Lines", "Received", "First Logged", "Last Changed", "Last Change", "Notes"])
    logged = datetime(2026, 9, 20, 8, 30)
    log.append(["SO-1328", "GNB Doors", "PO-000999", "UPWARDOR", "Open", date(2026, 9, 15),
                3, "1/3", logged, logged, "PO linked (Open)", "plastic panels only"])

    buf = io.BytesIO()
    wb.save(buf)
    content = buf.getvalue()

    parsed = svc.parse_records_from_bytes(content)
    assert parsed["SO-1328"].fulfillment == "Buy Complete"
    assert parsed["SO-1328"].fulfillment_explicit is False
    assert parsed["SO-1308"].fulfillment == "Emergency Build"
    assert parsed["SO-1308"].fulfillment_explicit is True
    assert parsed["SO-1308"].emergency_build == ""

    prior_log = svc.parse_po_log_from_bytes(content)
    assert prior_log[("SO-1328", "PO-000999")]["Notes"] == "plastic panels only"

    orders = [
        _so("SO-1328", [
            ("PN80-21100-1002", 4, "SECTION, PANORAMA"),
            ("PN45-24400-1400", 8),
        ], customer="GNB Doors"),
        _so("SO-1308", [
            ("PN80-21100-0802", 2, "SECTION, PANORAMA"),
            ("HK02-1", 1),
        ], customer="Silver Spray"),
        _so("SO-1299", [("PN20-2100010-0802", 3, "SECTION, SOLALITE")], customer="New panorama"),
    ]
    po = {"SO-1328": [_po_line("PN45-24400-1400", 8, po="PO-000999")]}
    rows, data, wb_out = _schedule(
        orders, records=parsed, po_lines_by_so=po, po_log_prior=prior_log, now=datetime(2026, 10, 8, 12, 30),
    )
    by_so = {r["SO Number"]: r for r in rows}
    assert by_so["SO-1328"]["Fulfillment"] == "Partial Build"
    assert by_so["SO-1328"][INHOUSE_BUILD_HEADER] == "Not Started"
    assert by_so["SO-1328"]["Order Status"] == "Shipped by Vendor"
    assert by_so["SO-1328"]["PO Date"].date() == date(2026, 9, 15)
    assert by_so["SO-1308"]["Fulfillment"] == "Emergency Build"
    assert by_so["SO-1308"]["Emergency Build"] == "Not Started"
    assert by_so["SO-1308"][INHOUSE_BUILD_HEADER] is None
    assert by_so["SO-1308"]["Shipping Status"] == "Ready to Ship"
    assert by_so["SO-1308"]["Pull from stock"] == "Yes"
    assert by_so["SO-1299"]["Fulfillment"] == "Partial Build"

    log_ws = wb_out["SO-PO Log"]
    log_header = [c.value for c in log_ws[1]]
    log_rows = [dict(zip(log_header, r)) for r in log_ws.iter_rows(min_row=2, values_only=True)]
    kept = [r for r in log_rows if r["SO Number"] == "SO-1328" and r["PO Number"] == "PO-000999"]
    assert len(kept) == 1
    assert kept[0]["Notes"] == "plastic panels only"

    # A person then sets that Partial Build back to Buy Complete. The file
    # now has the In-house Build column, so the next refresh keeps it.
    schedule = wb_out["Schedule"]
    header = [c.value for c in schedule[1]]
    ful_col = header.index("Fulfillment") + 1
    for row in schedule.iter_rows(min_row=2):
        if row[0].value == "SO-1328":
            schedule.cell(row=row[0].row, column=ful_col, value="Buy Complete")
    edited = io.BytesIO()
    wb_out.save(edited)
    reread = svc.parse_records_from_bytes(edited.getvalue())
    assert reread["SO-1328"].fulfillment_explicit is True
    rows2, _, _ = _schedule(orders, records=reread, po_lines_by_so=po)
    assert {r["SO Number"]: r["Fulfillment"] for r in rows2}["SO-1328"] == "Buy Complete"


def test_fulfillment_dropdown_includes_partial_build():
    _, data, wb = _schedule([_so("SO-1", [("PN45-24400-1400", 1)])])
    formulas = [str(dv.formula1) for dv in wb["Schedule"].data_validations.dataValidation]
    assert '"Buy Complete,Partial Build,Emergency Build"' in formulas
    assert '"Not Started,In Production,Complete,"' in formulas
    assert [c.value for c in wb["Schedule"][1]] == SCHEDULE_HEADERS
    note = wb["Schedule"].cell(row=1, column=SCHEDULE_HEADERS.index("Fulfillment") + 1).comment
    assert note is not None and "Partial Build" in note.text
    assert data  # workbook bytes written
