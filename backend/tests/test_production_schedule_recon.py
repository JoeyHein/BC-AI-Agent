"""Schedule ack / Pull from stock / reconciliation.

Graph and Business Central are mocked. Nothing here downloads SharePoint or
writes a purchase order.
"""
import io
import os
import sys
from datetime import date, datetime

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from app.config import settings
from app.services.production_schedule_service import (
    PO_LINKS_HEADERS,
    SCHEDULE_HEADERS,
    production_schedule_service as svc,
    partition_purchase_orders,
    _SORecord,
)
from app.services.upwardor_reconciliation import (
    ACK_HEADER,
    FLAG_HEADER,
    FLAG_PULL,
    NOTE_HEADER,
    NOTE_MISSED,
    NOTE_RED,
    PULL_HEADER,
    RECON_SHEET_NAME,
    SCHEDULE_RECON_HEADERS,
    STATUS_HEADER,
    parse_upwardor_open_order_export,
)

ORIGINAL_HEADERS = [
    "SO Number", "Customer Name", "Customer Tag / External Doc #", "Order Date", "PO Date",
    "Fulfillment", "Upwardor PO #", "Order Status", "Expected Receipt",
    "Operators", "Window Kits", "Emergency Build", "Shipping Status",
]


def _so(number, items, customer="Acme Doors"):
    return {
        "number": number, "customerName": customer, "externalDocumentNumber": "TAG",
        "orderDate": "2026-09-30",
        "salesOrderLines": [
            {"lineType": "Item", "lineObjectNumber": item, "quantity": qty, "description": item}
            for item, qty in items
        ],
    }


def _po_line(item, qty, po="PO-001", status="Open", received=0, vendor="UPWARDOR"):
    return {
        "lineType": "Item", "lineObjectNumber": item, "quantity": qty,
        "receivedQuantity": received, "expectedReceiptDate": "2026-10-15",
        "_po_number": po, "_po_status": status, "_vendor": vendor, "_po_date": "2026-10-01",
    }


def _ack(number, po, status="confirmed", received_at=None):
    return {
        "vendor_order_no": number,
        "our_po_number": po,
        "status": status,
        "completion_date": date(2026, 11, 20),
        "received_at": received_at or datetime(2026, 10, 1, 15, 0),
    }


def _build(orders, records=None, po_lines_by_so=None, vendor_acks=None, upwardor_open_orders=None,
           unlinked_pos=None):
    facts = svc.compute_so_facts(orders, po_lines_by_so or {})
    data, _, _ = svc.build_workbook_bytes(
        orders, records or {}, facts,
        po_lines_by_so=po_lines_by_so,
        vendor_acks=vendor_acks,
        upwardor_open_orders=upwardor_open_orders,
        unlinked_pos=unlinked_pos,
    )
    wb = load_workbook(io.BytesIO(data))
    ws = wb["Schedule"]
    header = [c.value for c in ws[1]]
    rows = [dict(zip(header, row)) for row in ws.iter_rows(min_row=2, values_only=True)]
    return rows, wb, data


def _rules(ws):
    found = []
    for sqref, rules in ws.conditional_formatting._cf_rules.items():
        for rule in rules:
            formula = rule.formula
            if formula and not isinstance(formula, str):
                formula = formula[0] if formula else ""
            found.append((str(sqref), str(formula), rule))
    return found


# ── columns stay, new ones append ───────────────────────────────────────

def test_existing_schedule_and_po_headers_stay_in_place():
    assert SCHEDULE_HEADERS[:len(ORIGINAL_HEADERS)] == ORIGINAL_HEADERS
    assert SCHEDULE_HEADERS[len(ORIGINAL_HEADERS):] == SCHEDULE_RECON_HEADERS
    assert PO_LINKS_HEADERS == [
        "SO Number", "Customer", "PO Number", "Vendor", "PO Status",
        "Lines", "Received", "Expected Receipt",
        "Vendor Ack #", "Ack Status", "Ack Received", "Completion Date",
    ]


def test_existing_dropdowns_remain_and_pull_from_stock_is_yes_blank():
    import zipfile
    rows, wb, data = _build([_so("SO-1", [("PN45-24400-1400", 4)])])
    formulas = [str(dv.formula1) for dv in wb["Schedule"].data_validations.dataValidation]
    assert '"Buy Complete,Emergency Build"' in formulas
    assert '"Waiting to Order,PO Drafted,Ordered,Shipped by Vendor,Partially Received,Received,"' in formulas
    assert '"Not Started,In Progress,Complete,"' in formulas
    assert '"Not Started,In Production,Complete,"' in formulas
    assert '"Not Ready,Ready to Ship,Shipped"' in formulas
    pull = get_column_letter(SCHEDULE_HEADERS.index(PULL_HEADER) + 1)
    yes = [dv for dv in wb["Schedule"].data_validations.dataValidation if str(dv.sqref).startswith(pull)]
    assert len(yes) == 1
    assert yes[0].formula1 == '"Yes"'
    # Blank cells are valid. openpyxl's reader reports allow_blank False; the file says 1.
    import re
    xml = zipfile.ZipFile(io.BytesIO(data)).read("xl/worksheets/sheet1.xml").decode()
    pull_dv = re.search(r'<dataValidation sqref="R2"[^>]*>', xml)
    assert pull_dv and 'allowBlank="1"' in pull_dv.group(0)
    assert '<formula1>"Yes"</formula1>' in xml
    assert rows[0][PULL_HEADER] is None


# ── flags ───────────────────────────────────────────────────────────────

def test_po_without_ack_is_red_and_purchase_order_row_is_filled():
    orders = [_so("SO-1", [("PN45-24400-1400", 4)])]
    po = {"SO-1": [_po_line("PN45-24400-1400", 4, po="PO-000958")]}
    rows, wb, _ = _build(orders, po_lines_by_so=po, vendor_acks=[])
    assert rows[0]["Upwardor PO #"] == "PO-000958"
    assert rows[0][ACK_HEADER] is None
    assert rows[0][STATUS_HEADER] == "Not acknowledged"
    assert rows[0][NOTE_HEADER] == NOTE_RED
    assert rows[0][FLAG_HEADER] == '=IF($R2="Yes","OK - pull from stock","RED")'
    formulas = [formula for _, formula, _ in _rules(wb["Schedule"])]
    assert 'AND($R2<>"Yes",$P2="RED")' in formulas
    red = next(rule for _, formula, rule in _rules(wb["Schedule"]) if "RED" in formula and "AND" in formula)
    assert red.dxf.font.bold is True
    assert red.dxf.font.italic is not True
    assert "FFC7CE" in str(red.dxf.fill.fgColor.rgb)
    assert "9C0006" in str(red.dxf.font.color.rgb)
    po_row = list(wb["Purchase Orders"].iter_rows(min_row=2, max_row=2))[0]
    assert po_row[2].value == "PO-000958"
    assert "FFC7CE" in str(po_row[0].fill.fgColor.rgb)


def test_open_so_with_no_po_and_no_ack_is_missed():
    rows, wb, _ = _build([_so("SO-1", [("PN45-24400-1400", 4)])])
    assert rows[0][STATUS_HEADER] == "No Upwardor PO, no ack"
    assert rows[0][NOTE_HEADER] == NOTE_MISSED
    assert "MISSED" in rows[0][FLAG_HEADER]
    formulas = [formula for _, formula, _ in _rules(wb["Schedule"])]
    assert 'AND($R2<>"Yes",$P2="MISSED")' in formulas
    missed = next(rule for _, formula, rule in _rules(wb["Schedule"]) if "MISSED" in formula)
    assert missed.dxf.font.bold is True
    assert missed.dxf.font.italic is True
    assert "F8CBAD" in str(missed.dxf.fill.fgColor.rgb)
    assert "833C0B" in str(missed.dxf.font.color.rgb)
    # Distinct from the red rule, and both sit above the column colors (priority 1 and 2).
    listed = [(str(sqref), formula) for sqref, formula, _ in _rules(wb["Schedule"])]
    assert "A2:R" in listed[0][0] and "RED" in listed[0][1]
    assert "A2:R" in listed[1][0] and "MISSED" in listed[1][1]


def test_operator_only_and_wrap_only_are_not_missed():
    rows, _, _ = _build([
        _so("SO-1", [("OP20-00001", 1)]),
        _so("SO-2", [("WRAP-001", 1)]),
    ])
    for row in rows:
        assert row[NOTE_HEADER] is None
        assert row[FLAG_HEADER].endswith(',"")')
        assert row["Order Status"] is None


def test_acknowledged_po_shows_sord_and_is_ok():
    orders = [_so("SO-1", [("PN45-24400-1400", 4)])]
    po = {"SO-1": [_po_line("PN45-24400-1400", 4, po="PO-000956")]}
    acks = [_ack("S-ORD116307", "PO-000956")]
    rows, wb, _ = _build(orders, po_lines_by_so=po, vendor_acks=acks)
    assert rows[0][ACK_HEADER] == "S-ORD116307"
    assert rows[0][STATUS_HEADER] == "Confirmed"
    assert rows[0][FLAG_HEADER].endswith(',"OK")')
    po_row = list(wb["Purchase Orders"].iter_rows(min_row=2, max_row=2))[0]
    assert po_row[8].value == "S-ORD116307"
    assert po_row[0].fill.patternType is None


def test_one_unacked_po_among_two_stays_red_and_keeps_the_ack_we_have():
    orders = [_so("SO-1", [("PN45-1", 1), ("PN45-2", 1)])]
    po = {"SO-1": [
        _po_line("PN45-1", 1, po="PO-000956"),
        _po_line("PN45-2", 1, po="PO-000959"),
    ]}
    rows, _, _ = _build(orders, po_lines_by_so=po, vendor_acks=[_ack("S-ORD116307", "PO-000956")])
    assert "S-ORD116307" in rows[0][ACK_HEADER]
    assert "PO-000959" in rows[0][NOTE_HEADER]
    assert rows[0][NOTE_HEADER].startswith(NOTE_RED)
    assert rows[0][FLAG_HEADER].endswith(',"RED")')


def test_cancelled_ack_is_check_not_red():
    orders = [_so("SO-1", [("PN45-1", 1)])]
    po = {"SO-1": [_po_line("PN45-1", 1, po="PO-000958")]}
    rows, wb, _ = _build(orders, po_lines_by_so=po, vendor_acks=[_ack("S-ORD116000", "PO-000958", status="cancelled")])
    assert rows[0][ACK_HEADER] == "S-ORD116000"
    assert rows[0][STATUS_HEADER] == "Cancelled"
    assert rows[0][FLAG_HEADER].endswith(',"CHECK")')
    assert "cancelled" in rows[0][NOTE_HEADER].lower()
    formulas = [formula for _, formula, _ in _rules(wb["Schedule"])]
    assert '$P2="CHECK"' in formulas


def test_pull_from_stock_yes_is_not_counted_as_missed_and_formula_suppresses_it():
    prior = _SORecord("SO-1")
    prior.pull_from_stock = "Yes"
    rows, _, _ = _build([_so("SO-1", [("PN45-1", 1)])], records={"SO-1": prior})
    assert rows[0][PULL_HEADER] == "Yes"
    assert rows[0][FLAG_HEADER] == f'=IF($R2="Yes","{FLAG_PULL}","MISSED")'
    assert rows[0][NOTE_HEADER] == NOTE_MISSED
    assert svc.last_recon["pull_from_stock"] == 1
    assert svc.last_recon["missed"] == 0
    assert svc.last_recon["red"] == 0


def test_pull_from_stock_round_trips_by_header_name_not_position():
    _, _, data = _build([_so("SO-1", [("PN45-1", 1)])])
    wb = load_workbook(io.BytesIO(data))
    header = [c.value for c in wb["Schedule"][1]]
    wb["Schedule"].cell(row=2, column=header.index(PULL_HEADER) + 1, value="Yes")
    buf = io.BytesIO()
    wb.save(buf)

    shuffled = Workbook()
    ws = shuffled.active
    ws.title = "Schedule"
    ws.append([PULL_HEADER, "Shipping Status", "SO Number", "Fulfillment"])
    ws.append(["Yes", "Shipped", "SO-1", "Emergency Build"])
    shuffled_buf = io.BytesIO()
    shuffled.save(shuffled_buf)

    from_rebuild = svc.parse_records_from_bytes(buf.getvalue())["SO-1"]
    from_shuffle = svc.parse_records_from_bytes(shuffled_buf.getvalue())["SO-1"]
    assert from_rebuild.pull_from_stock == "Yes"
    assert from_rebuild.recon_flag == "MISSED"
    assert from_shuffle.pull_from_stock == "Yes"
    assert from_shuffle.fulfillment == "Emergency Build"
    assert from_shuffle.shipping_status == "Shipped"

    rows, _, _ = _build(
        [_so("SO-1", [("PN45-1", 1)])],
        records={"SO-1": from_rebuild},
    )
    assert rows[0][PULL_HEADER] == "Yes"
    assert rows[0][FLAG_HEADER].endswith(',"MISSED")')


def test_archived_flag_formula_uses_the_archived_row():
    """A formula read from schedule row 3 must not be replayed as $R3 on archive row 2."""
    orders = [_so("SO-1", [("PN45-1", 1)]), _so("SO-2", [("PN45-1", 1)])]
    _, _, data = _build(orders)
    wb = load_workbook(io.BytesIO(data))
    header = [c.value for c in wb["Schedule"][1]]
    wb["Schedule"].cell(row=3, column=header.index(PULL_HEADER) + 1, value="Yes")
    buf = io.BytesIO()
    wb.save(buf)
    records = svc.parse_records_from_bytes(buf.getvalue())
    assert records["SO-2"].pull_from_stock == "Yes"
    assert records["SO-2"].recon_flag == "MISSED"

    _, wb2, _ = _build([_so("SO-1", [("PN45-1", 1)])], records=records)
    archived = list(wb2["Archived"].iter_rows(min_row=2, values_only=True))
    assert archived[0][0] == "SO-2"
    flag = archived[0][SCHEDULE_HEADERS.index(FLAG_HEADER)]
    pull = archived[0][SCHEDULE_HEADERS.index(PULL_HEADER)]
    assert pull == "Yes"
    assert flag == f'=IF($R2="Yes","{FLAG_PULL}","MISSED")'


def test_open_order_export_annotates_status_and_flags_a_missing_sord():
    orders = [
        _so("SO-1", [("PN45-24400-1400", 4)]),
        _so("SO-2", [("PN45-24400-1400", 2)]),
    ]
    po = {
        "SO-1": [_po_line("PN45-24400-1400", 4, po="PO-000956")],
        "SO-2": [_po_line("PN45-24400-1400", 2, po="PO-000866")],
    }
    acks = [
        _ack("S-ORD116307", "PO-000956"),
        _ack("S-ORD115523", "PO-000866"),
    ]
    export = [{
        "sord": "S-ORD116307",
        "customer": "Acme",
        "line_count": 1,
        "sections_ordered": 4,
        "sections_remaining": 4,
        "sections_to_mfg": 4,
        "lines": [],
    }]
    rows, _, _ = _build(orders, po_lines_by_so=po, vendor_acks=acks, upwardor_open_orders=export)
    by_so = {row["SO Number"]: row for row in rows}
    assert "Open on Upwardor list" in by_so["SO-1"][STATUS_HEADER]
    assert "4 of 4 remaining" in by_so["SO-1"][STATUS_HEADER]
    assert by_so["SO-1"][FLAG_HEADER].endswith(',"OK")')
    assert by_so["SO-2"][FLAG_HEADER].endswith(',"CHECK")')
    assert "Not on Upwardor open-order list" in (by_so["SO-2"][STATUS_HEADER] or "")


# ── reconciliation sheet + export parser ────────────────────────────────

def _export_bytes(rows, header_row=1):
    wb = Workbook()
    ws = wb.active
    ws.title = "Upw Sales Order Master"
    headers = [
        "Customer Name", "Sales Order No.", "Part No.", "Description",
        "Qty. Ordered", "remaining Qty", "Qty per UOM", "Qty. to Manufacture",
        "print UOM", "Bulk",
    ]
    if header_row > 1:
        ws.append(["Open orders"])
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parser_groups_sords_and_counts_only_sec_lines():
    content = _export_bytes([
        ["OPEN DISTRIBUTION COMPANY INC.", "S-ORD115492", "PN65-24400-1600", "SECTION", 1, 1, 1, 0, "EA", "SEC"],
        ["OPEN DISTRIBUTION COMPANY INC.", "S-ORD115492", "PL10-00146-02", "RETAINER", 1, 1, 1, 0, "EA", "Bulk"],
        ["OPEN DISTRIBUTION COMPANY INC.", "S-ORD115491", "FH17-00018-00", "STRUT", 10, 10, 1, 0, "EA", "Bulk"],
    ], header_row=2)
    parsed = parse_upwardor_open_order_export(content)
    by_sord = {row["sord"]: row for row in parsed}
    assert set(by_sord) == {"S-ORD115492", "S-ORD115491"}
    assert by_sord["S-ORD115492"]["line_count"] == 2
    assert by_sord["S-ORD115492"]["sections_ordered"] == 1
    assert by_sord["S-ORD115492"]["lines"][0]["part_no"] == "PN65-24400-1600"
    assert by_sord["S-ORD115491"]["sections_ordered"] == 0


def test_parser_rejects_a_file_that_is_not_the_export():
    wb = Workbook()
    wb.active.append(["Hello"])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(ValueError, match="Sales Order No."):
        parse_upwardor_open_order_export(buf.getvalue())


def test_reconciliation_sheet_without_export_lists_unacked_and_stock_pos():
    orders = [_so("SO-1", [("PN45-1", 4)])]
    po = {"SO-1": [_po_line("PN45-1", 4, po="PO-000958")]}
    unlinked = [
        {"po_number": "PO-000797", "vendor": "UPWARDOR", "po_status": "Open", "lines": [
            _po_line("EC-1", 10, po="PO-000797"),
        ]},
        {"po_number": "PO-LIFT", "vendor": "LIFTMASTER", "po_status": "Open", "lines": [
            _po_line("OP20-1", 1, po="PO-LIFT"),
        ]},
    ]
    _, wb, _ = _build(orders, po_lines_by_so=po, unlinked_pos=unlinked)
    ws = wb[RECON_SHEET_NAME]
    assert "not loaded" in ws.cell(row=1, column=1).value
    values = list(ws.iter_rows(min_row=1, values_only=True))
    flat = [cell for row in values for cell in row if cell]
    assert "PO-000958" in flat
    assert "PO-000797" in flat
    assert "PO-LIFT" not in flat
    assert NOTE_RED in flat
    assert RECON_SHEET_NAME in wb.sheetnames
    assert wb.sheetnames == [
        "Schedule", "Archived", "Assignments", "Purchase Orders", "SO-PO Log", RECON_SHEET_NAME,
    ]


def test_reconciliation_sheet_compares_export_qty_and_unmatched_sords():
    orders = [_so("SO-1", [("PN65-24400-1600", 6)], customer="Band City")]
    po = {"SO-1": [_po_line("PN65-24400-1600", 6, po="PO-000967")]}
    acks = [_ack("S-ORD116302", "PO-000967")]
    content = _export_bytes([
        ["OPEN DISTRIBUTION COMPANY INC.", "S-ORD116302", "PN65-24400-1600", "SECTION", 8, 8, 1, 0, "EA", "SEC"],
        ["OPEN DISTRIBUTION COMPANY INC.", "S-ORD115687", "OP99-1", "OPERATOR", 2, 2, 1, 0, "EA", "Bulk"],
    ])
    export = parse_upwardor_open_order_export(content)
    _, wb, _ = _build(orders, po_lines_by_so=po, vendor_acks=acks, upwardor_open_orders=export)
    ws = wb[RECON_SHEET_NAME]
    text = [cell for row in ws.iter_rows(values_only=True) for cell in row if isinstance(cell, str)]
    joined = "\n".join(text)
    assert "S-ORD116302" in joined
    assert "Matched - qty mismatch" in joined
    assert "BC PO-000967 has 6 sections" in joined
    assert "S-ORD115687" in joined
    assert "No OpenDC PO" in joined
    assert "SO-1" in joined or "Band City" in joined


def test_partition_splits_linked_sales_orders_from_stock_pos():
    pos = [
        {
            "number": "PO-000958", "status": "Draft", "vendorName": "UPWARDOR",
            "orderDate": "2026-10-01", "externalDocumentNumber": "SO-001294",
            "purchaseOrderLines": [
                {"lineType": "Item", "lineObjectNumber": "PN45-1", "quantity": 4, "description": "section"},
            ],
        },
        {
            "number": "PO-000797", "status": "Open", "vendorName": "UPWARDOR",
            "orderDate": "2026-05-01", "externalDocumentNumber": "",
            "purchaseOrderLines": [
                {"lineType": "Item", "lineObjectNumber": "EC-1", "quantity": 10, "description": "end caps"},
            ],
        },
    ]
    by_so, unlinked = partition_purchase_orders(pos)
    assert by_so["SO-001294"][0]["_po_number"] == "PO-000958"
    assert by_so["SO-001294"][0]["_vendor"] == "UPWARDOR"
    assert [row["po_number"] for row in unlinked] == ["PO-000797"]


def test_build_and_deliver_reads_pull_from_stock_without_calling_graph_or_bc(monkeypatch):
    orders = [_so("SO-1", [("PN45-1", 1)])]
    rows, _, data = _build(orders)
    assert rows[0][PULL_HEADER] is None
    wb = load_workbook(io.BytesIO(data))
    header = [c.value for c in wb["Schedule"][1]]
    wb["Schedule"].cell(row=2, column=header.index(PULL_HEADER) + 1, value="Yes")
    buf = io.BytesIO()
    wb.save(buf)
    existing = buf.getvalue()

    calls = []

    def download(drive_id, path):
        calls.append(("download", drive_id, path))
        return existing

    uploaded = {}

    def upload(drive_id, path, content):
        calls.append(("upload", drive_id, path))
        uploaded["content"] = content
        return "https://example.invalid/schedule"

    def boom(*args, **kwargs):
        raise AssertionError("live Graph or BC was called")

    import app.services.production_schedule_service as mod
    monkeypatch.setattr(settings, "PRODSCHED_SHAREPOINT_ENABLED", True)
    monkeypatch.setattr(settings, "PRODSCHED_SHAREPOINT_DRIVE_ID", "drive-test")
    monkeypatch.setattr(settings, "PRODSCHED_SHAREPOINT_FILE_PATH", "Production Schedule/OPENDC_Production_Schedule.xlsx")
    monkeypatch.setattr(mod.graph_client, "download_drive_file", download)
    monkeypatch.setattr(mod.graph_client, "upload_drive_file", upload)
    monkeypatch.setattr(mod.graph_client, "_get_access_token", boom)
    monkeypatch.setattr(mod.bc_client, "get_open_sales_orders_with_lines", boom)
    monkeypatch.setattr(mod.bc_client, "get_open_purchase_orders_with_lines", boom)
    monkeypatch.setattr(svc, "_fetch_bc", lambda: {
        "orders": orders,
        "po_lines_by_so": {},
        "unlinked_pos": [],
        "so_facts": svc.compute_so_facts(orders, {}),
        "prod_orders": [],
        "prod_so_map": {},
        "picking_remaining": {},
        "vendor_acks": [],
        "upwardor_open_orders": None,
    })
    monkeypatch.setattr(svc, "_load_vendor_acks", boom)

    result = svc.build_and_deliver()
    assert calls == [
        ("download", "drive-test", "Production Schedule/OPENDC_Production_Schedule.xlsx"),
        ("upload", "drive-test", "Production Schedule/OPENDC_Production_Schedule.xlsx"),
    ]
    assert result["recon_pull_from_stock"] == 1
    assert result["recon_missed"] == 0
    assert result["sharepoint"] == "https://example.invalid/schedule"
    out = load_workbook(io.BytesIO(uploaded["content"]))
    out_header = [c.value for c in out["Schedule"][1]]
    assert out["Schedule"].cell(row=2, column=out_header.index(PULL_HEADER) + 1).value == "Yes"
    assert out_header[:len(ORIGINAL_HEADERS)] == ORIGINAL_HEADERS
    assert RECON_SHEET_NAME in out.sheetnames
