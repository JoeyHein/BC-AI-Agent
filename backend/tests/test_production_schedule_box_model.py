"""Schedule sheet under the box-in/box-out model (2026-10-02): one row per SO,
whole-order status from the Upwardor PO that mirrors it, operators tracked
separately, window kits as the only default in-house work, Emergency Build
as the exception. See production_schedule_service module docstring.
"""
import io
import os
import sys
from datetime import date

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from openpyxl import Workbook, load_workbook

from app.services.production_schedule_service import (
    production_schedule_service as svc,
    SCHEDULE_HEADERS,
    _SORecord,
)


def _so(number, items, customer="Acme Doors"):
    return {
        "number": number, "customerName": customer, "externalDocumentNumber": "TAG",
        "orderDate": "2026-09-30",
        "salesOrderLines": [
            {"lineType": "Item", "lineObjectNumber": item, "quantity": qty, "description": item}
            for item, qty in items
        ],
    }


def _po_line(item, qty, received=0, po="PO-001", status="Open", expected="2026-10-15"):
    return {"lineType": "Item", "lineObjectNumber": item, "quantity": qty,
            "receivedQuantity": received, "expectedReceiptDate": expected,
            "_po_number": po, "_po_status": status, "_vendor": "UPWARDOR"}


def _schedule(orders, records=None, po_lines_by_so=None):
    facts = svc.compute_so_facts(orders, po_lines_by_so or {})
    data, _, _ = svc.build_workbook_bytes(orders, records or {}, facts, po_lines_by_so=po_lines_by_so)
    ws = load_workbook(io.BytesIO(data))["Schedule"]
    header = [c.value for c in ws[1]]
    return [dict(zip(header, row)) for row in ws.iter_rows(min_row=2, values_only=True)], data


# ── whole-order status from linked POs ──────────────────────────────────

def test_no_po_yet_is_waiting_to_order():
    rows, _ = _schedule([_so("SO-1", [("PN45-24400-1400", 4)])])
    assert rows[0]["Order Status"] == "Waiting to Order"
    assert rows[0]["Fulfillment"] == "Buy Complete"
    assert rows[0]["Upwardor PO #"] is None


def test_draft_po_is_po_drafted_and_shows_po_number():
    rows, _ = _schedule([_so("SO-1", [("PN45-24400-1400", 4)])],
                        po_lines_by_so={"SO-1": [_po_line("PN45-24400-1400", 4, status="Draft")]})
    assert rows[0]["Order Status"] == "PO Drafted"
    assert rows[0]["Upwardor PO #"] == "PO-001"


def test_released_po_is_ordered_with_expected_receipt():
    rows, _ = _schedule([_so("SO-1", [("PN45-24400-1400", 4)])],
                        po_lines_by_so={"SO-1": [_po_line("PN45-24400-1400", 4)]})
    assert rows[0]["Order Status"] == "Ordered"
    assert rows[0]["Expected Receipt"].date() == date(2026, 10, 15)


def test_partial_and_full_receipt():
    po = {"SO-1": [_po_line("PN45-24400-1400", 4, received=4), _po_line("HK02-1", 1, received=0)]}
    rows, _ = _schedule([_so("SO-1", [("PN45-24400-1400", 4), ("HK02-1", 1)])], po_lines_by_so=po)
    assert rows[0]["Order Status"] == "Partially Received"

    po = {"SO-1": [_po_line("PN45-24400-1400", 4, received=4), _po_line("HK02-1", 1, received=1)]}
    rows, _ = _schedule([_so("SO-1", [("PN45-24400-1400", 4), ("HK02-1", 1)])], po_lines_by_so=po)
    assert rows[0]["Order Status"] == "Received"
    assert rows[0]["Expected Receipt"] is None


def test_status_never_moves_backward_from_hand_set_shipped_by_vendor():
    prior = _SORecord("SO-1")
    prior.order_status = "Shipped by Vendor"
    rows, _ = _schedule([_so("SO-1", [("PN45-24400-1400", 4)])], records={"SO-1": prior},
                        po_lines_by_so={"SO-1": [_po_line("PN45-24400-1400", 4)]})
    assert rows[0]["Order Status"] == "Shipped by Vendor"


def test_status_advances_past_hand_set_value_once_bc_catches_up():
    prior = _SORecord("SO-1")
    prior.order_status = "Shipped by Vendor"
    rows, _ = _schedule([_so("SO-1", [("PN45-24400-1400", 4)])], records={"SO-1": prior},
                        po_lines_by_so={"SO-1": [_po_line("PN45-24400-1400", 4, received=4)]})
    assert rows[0]["Order Status"] == "Received"


# ── operators / window kits / emergency ───────────────────────────────

def test_operators_tracked_separately_from_upwardor_po():
    po = {"SO-1": [_po_line("PN45-24400-1400", 4, received=4)]}
    rows, _ = _schedule([_so("SO-1", [("PN45-24400-1400", 4), ("OP20-00001", 1)])], po_lines_by_so=po)
    assert rows[0]["Order Status"] == "Received"
    assert rows[0]["Operators"] == "Waiting to Order"


def test_no_operator_on_so_leaves_operators_blank():
    rows, _ = _schedule([_so("SO-1", [("PN45-24400-1400", 4)])])
    assert rows[0]["Operators"] is None


def test_window_kits_seeded_only_when_so_has_gk_items():
    rows, _ = _schedule([_so("SO-1", [("PN45-24400-1400", 4), ("GK17-10000-00", 2)]),
                         _so("SO-2", [("PN45-24400-1400", 4)])])
    assert rows[0]["Window Kits"] == "Not Started"
    assert rows[1]["Window Kits"] is None


def test_hand_edited_window_kit_status_carries_forward():
    prior = _SORecord("SO-1")
    prior.window_kits = "Complete"
    rows, _ = _schedule([_so("SO-1", [("GK17-10000-00", 2)])], records={"SO-1": prior})
    assert rows[0]["Window Kits"] == "Complete"


def test_emergency_build_seeded_when_flipped_and_cleared_when_flipped_back():
    prior = _SORecord("SO-1")
    prior.fulfillment = "Emergency Build"
    rows, _ = _schedule([_so("SO-1", [("PN45-24400-1400", 4)])], records={"SO-1": prior})
    assert rows[0]["Emergency Build"] == "Not Started"

    prior.fulfillment = "Buy Complete"
    prior.emergency_build = "In Production"
    rows, _ = _schedule([_so("SO-1", [("PN45-24400-1400", 4)])], records={"SO-1": prior})
    assert rows[0]["Emergency Build"] is None


def test_wrap_only_so_has_nothing_for_upwardor():
    rows, _ = _schedule([_so("SO-1", [("WRAP-001", 1), ("OP20-00001", 1)])])
    assert rows[0]["Order Status"] is None


# ── read-back ─────────────────────────────────────────────────────────

def test_round_trip_preserves_hand_edits():
    prior = _SORecord("SO-1")
    prior.fulfillment = "Emergency Build"
    prior.emergency_build = "In Production"
    prior.window_kits = "In Progress"
    prior.shipping_status = "Ready to Ship"
    prior.po_date = date(2026, 9, 1)
    _, data = _schedule([_so("SO-1", [("GK17-10000-00", 2)])], records={"SO-1": prior})

    rec = svc.parse_records_from_bytes(data)["SO-1"]
    assert rec.fulfillment == "Emergency Build"
    assert rec.emergency_build == "In Production"
    assert rec.window_kits == "In Progress"
    assert rec.shipping_status == "Ready to Ship"
    assert rec.po_date == date(2026, 9, 1)


def test_read_back_is_by_header_name_not_position():
    """A sheet written with columns in a different order / missing a column
    must still parse — guards the 2026-08-25 silent-wipe failure mode."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Schedule"
    ws.append(["SO Number", "Shipping Status", "Fulfillment"])
    ws.append(["SO-1", "Shipped", "Emergency Build"])
    buf = io.BytesIO()
    wb.save(buf)

    rec = svc.parse_records_from_bytes(buf.getvalue())["SO-1"]
    assert rec.shipping_status == "Shipped"
    assert rec.fulfillment == "Emergency Build"


def test_migrates_v3_component_layout():
    """The live file before 2026-10-02: 2 header rows, 7 components ×
    Purchasing/Production. Shipping + PO date carry; per-component
    production is dropped; a hand-set Shipped by Vendor carries into Order
    Status; stock-netting Received does NOT (meant raw parts on hand)."""
    components = ["Panels", "Hardware", "Tracks", "Springs", "Shafts", "Weather Stripping", "Operators"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Schedule"
    row1 = ["SO Number", "Customer Name", "Customer Tag / External Doc #", "Order Date", "PO Date"]
    row2 = [None] * 5
    for c in components:
        row1 += [c, None]
        row2 += ["Purchasing", "Production"]
    ws.append(row1 + ["Shipping Status"])
    ws.append(row2 + [None])

    def data_row(so, panels_purch, ops_purch):
        row = [so, "Acme", "TAG", date(2026, 9, 1), date(2026, 9, 2)]
        for c in components:
            purch = {"Panels": panels_purch, "Operators": ops_purch}.get(c, "Received")
            row += [purch, "Complete"]
        return row + ["Ready to Ship"]

    ws.append(data_row("SO-1", "Shipped by Vendor", "Waiting to Order"))
    ws.append(data_row("SO-2", "Received", "In Stock"))
    buf = io.BytesIO()
    wb.save(buf)

    recs = svc.parse_records_from_bytes(buf.getvalue())
    assert recs["SO-1"].order_status == "Shipped by Vendor"
    assert recs["SO-1"].operators == ""
    assert recs["SO-1"].shipping_status == "Ready to Ship"
    assert recs["SO-1"].po_date == date(2026, 9, 2)
    assert recs["SO-2"].order_status == "Waiting to Order"
    assert recs["SO-2"].operators == "Received"
    assert recs["SO-2"].fulfillment == "Buy Complete"


def test_closed_so_moves_to_archived():
    prior = _SORecord("SO-9")
    prior.shipping_status = "Shipped"
    _, data = _schedule([_so("SO-1", [("PN45-24400-1400", 4)])], records={"SO-9": prior})
    wb = load_workbook(io.BytesIO(data))
    archived = [r[0] for r in wb["Archived"].iter_rows(min_row=2, values_only=True)]
    assert archived == ["SO-9"]
    assert [c.value for c in wb["Schedule"][1]] == SCHEDULE_HEADERS


def test_purchase_orders_sheet_lists_linked_pos_for_open_sos():
    po = {
        "SO-1": [_po_line("PN45-24400-1400", 4, received=4), _po_line("HK02-1", 1)],
        "SO-9": [_po_line("PN45-24400-1400", 1, po="PO-009")],  # SO not open
    }
    _, data = _schedule([_so("SO-1", [("PN45-24400-1400", 4), ("HK02-1", 1)])], po_lines_by_so=po)
    rows = list(load_workbook(io.BytesIO(data))["Purchase Orders"].iter_rows(min_row=2, values_only=True))
    assert len(rows) == 1
    assert rows[0][0] == "SO-1"
    assert rows[0][2] == "PO-001"
    assert rows[0][6] == "1/2"
