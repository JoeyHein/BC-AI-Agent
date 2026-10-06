"""Correlate a sales order, its purchase order(s), and Upwardor's acknowledgement.

The production schedule calls `schedule_recon` for every open SO. Ack numbers
and status come from vendor-ack intake (`vendor_order_acks` via
`ack_cells_for_po`) when a row exists. Pull from stock is not decided here —
the sheet writes a formula so a hand-set Yes clears the flag immediately, and
the next rebuild reads that Yes back by header name.

Upwardor's open-order export (their "Upw Sales Order Master" workbook, no PO
column) is optional. `parse_upwardor_open_order_export` reads it.
`build_reconciliation_report` compares those S-ORDs with our POs, matching
through the acknowledgement (S-ORD → PO). Nothing in this module downloads
mail or writes to SharePoint or Business Central. A later job can parse an
attachment and pass the rows in as `upwardor_open_orders`.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from typing import Dict, Iterable, List, Optional, Sequence

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.services.vendor_ack_parser import ack_cells_for_po

# Schedule columns appended after the box-in/box-out headers.
ACK_HEADER = "Upwardor Ack # (S-ORD)"
STATUS_HEADER = "Upwardor Status"
FLAG_HEADER = "Recon Flag"
NOTE_HEADER = "Recon Note"
PULL_HEADER = "Pull from stock"
SCHEDULE_RECON_HEADERS = [ACK_HEADER, STATUS_HEADER, FLAG_HEADER, NOTE_HEADER, PULL_HEADER]

PULL_FROM_STOCK_YES = "Yes"

FLAG_RED = "RED"
FLAG_MISSED = "MISSED"
FLAG_CHECK = "CHECK"
FLAG_OK = "OK"
FLAG_PULL = "OK - pull from stock"

NOTE_MISSED = "MISSED, no PO, no ack"
NOTE_RED = "NOT ACKNOWLEDGED by Upwardor"
NOTE_CANCELLED = "Upwardor acknowledgement was cancelled"
NOTE_NOT_ON_LIST = (
    "Not on the Upwardor open-order list — confirm whether it shipped "
    "or should be cancelled in BC"
)

RECON_SHEET_NAME = "Upwardor Reconciliation"

# Upw Sales Order Master export. Names match the workbook Upwardor sends
# (header row is detected; it is not assumed to be row 1).
_EXPORT_SALES_ORDER = "sales order no."
_EXPORT_PART = "part no."
_EXPORT_DESCRIPTION = "description"
_EXPORT_QTY_ORDERED = "qty. ordered"
_EXPORT_QTY_REMAINING = "remaining qty"
_EXPORT_QTY_TO_MFG = "qty. to manufacture"
_EXPORT_BULK = "bulk"
_EXPORT_CUSTOMER = "customer name"
_EXPORT_REQUIRED = (
    _EXPORT_SALES_ORDER,
    _EXPORT_QTY_ORDERED,
    _EXPORT_QTY_REMAINING,
    _EXPORT_QTY_TO_MFG,
    _EXPORT_BULK,
)

# BC section items. Used only to compare quantity with the export's SEC lines.
_BC_SECTION_PREFIXES = ("PN",)

_SORD_RE = re.compile(r"S-ORD\d+", re.IGNORECASE)

OPEN_ORDER_HEADERS = [
    "S-ORD (Ack #)", "Matched OpenDC PO", "Match Source", "SO (from PO)", "Customer",
    "Upwardor Lines", "Sections Ordered", "Sections Remaining", "Sections To Mfg",
    "BC PO Sections", "Result", "Note",
]
OUR_PO_HEADERS = [
    "PO Number", "Vendor Ack #", "SO Number", "Customer", "PO Status",
    "Recon Flag", "Note",
]

_RESULT_MATCHED = "Matched"
_RESULT_QTY = "Matched - qty mismatch"
_RESULT_NO_PO = "No OpenDC PO"
_RESULT_CLOSED = "Mismatch - closed/cancelled in BC"

_HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_SECTION_FILL = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
_SECTION_FONT = Font(color="1F4E78", bold=True)
_RED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
_RED_FONT = Font(color="9C0006", bold=True)
_AMBER_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
_AMBER_FONT = Font(color="9C6500")
_ORANGE_FILL = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")
_ORANGE_FONT = Font(color="833C0B", bold=True, italic=True)
_NOTE_FONT = Font(italic=True, color="595959")


@dataclass
class ScheduleRecon:
    """Computed columns for one schedule row, before the Pull-from-stock formula."""

    ack_numbers: Optional[str]
    upwardor_status: Optional[str]
    flag: str
    note: Optional[str]


@dataclass
class ReconciliationReport:
    """Rows for the Upwardor Reconciliation sheet. Read-only; rebuilt every run."""

    export_loaded: bool
    source_note: str
    open_order_rows: List[dict] = field(default_factory=list)
    our_po_rows: List[dict] = field(default_factory=list)
    export_count: int = 0


def normalize_pull_from_stock(value) -> str:
    """Yes survives a rebuild. Anything else (blank, typos) is blank."""
    if value is None:
        return ""
    return PULL_FROM_STOCK_YES if str(value).strip() == PULL_FROM_STOCK_YES else ""


_FLAG_FORMULA_RE = re.compile(
    r'^=IF\(\$[A-Z]+\d+="' + re.escape(PULL_FROM_STOCK_YES) + r'","'
    + re.escape(FLAG_PULL) + r'","(.*)"\)$'
)


def base_flag_from_cell(value) -> str:
    """Pull the computed flag back out of the sheet formula.

    The formula is rewritten on every row, so a stored `=IF($R5=...)` must
    not be copied onto a different row (archived SOs move). The else-branch
    is the flag the last rebuild computed.
    """
    if value is None:
        return ""
    text = str(value).strip()
    match = _FLAG_FORMULA_RE.match(text)
    if match:
        return match.group(1).replace('""', '"')
    if text.startswith("="):
        return ""
    return text


def recon_flag_formula(row: int, base_flag: str, pull_col: str) -> str:
    """Excel formula: Yes in Pull from stock wins over the computed flag.

    The else-branch keeps RED / MISSED / CHECK so clearing Yes on the next
    edit (or the next rebuild, which reads the cell back) shows the problem
    again without waiting on a new acknowledgement.
    """
    flag = (base_flag or "").replace('"', '""')
    return f'=IF(${pull_col}{row}="{PULL_FROM_STOCK_YES}","{FLAG_PULL}","{flag}")'


def recon_row_formulas(first_row: int, flag_col: str, pull_col: str) -> Dict[str, str]:
    """Conditional-format formulas anchored at the first data row.

    Pull from stock = Yes suppresses RED and MISSED. CHECK stays on the flag
    cell only — it is a review item, not a missing order.
    """
    return {
        FLAG_RED: f'AND(${pull_col}{first_row}<>"{PULL_FROM_STOCK_YES}",${flag_col}{first_row}="{FLAG_RED}")',
        FLAG_MISSED: f'AND(${pull_col}{first_row}<>"{PULL_FROM_STOCK_YES}",${flag_col}{first_row}="{FLAG_MISSED}")',
        FLAG_CHECK: f'${flag_col}{first_row}="{FLAG_CHECK}"',
    }


def sord_numbers(text: Optional[str]) -> List[str]:
    """S-ORD tokens in an ack cell, uppercased, first-seen order."""
    found: List[str] = []
    for match in _SORD_RE.findall(str(text or "")):
        token = match.upper()
        if token not in found:
            found.append(token)
    return found


def split_po_numbers(value: Optional[str]) -> List[str]:
    return [part.strip() for part in str(value or "").split(",") if part.strip()]


def schedule_recon(
    po_numbers: Optional[str],
    vendor_acks: Optional[Sequence[dict]],
    needs_upwardor: bool,
    open_orders_by_sord: Optional[Dict[str, dict]] = None,
) -> ScheduleRecon:
    """Flag one open sales order.

    RED — at least one PO in `po_numbers` has no acknowledgement number.
    MISSED — the SO needs an Upwardor buy, and it has no PO and no ack.
    CHECK — every PO is acknowledged, but the ack was cancelled, or an
            optional open-order export does not list the S-ORD.
    OK — acknowledged (or the export still shows the S-ORD open).
    Blank — nothing on the SO is bought from Upwardor (operator-only, wrap-only).

    `open_orders_by_sord` is None when no export was provided. An empty dict
    means the export was loaded and listed nothing.
    """
    pos = split_po_numbers(po_numbers)
    acks = list(vendor_acks or [])
    ack_labels: List[str] = []
    status_labels: List[str] = []
    missing: List[str] = []
    for po in pos:
        cells = ack_cells_for_po(po, acks)
        if cells.get("vendor_ack"):
            ack_labels.append(cells["vendor_ack"])
            if cells.get("ack_status"):
                status_labels.append(cells["ack_status"])
        else:
            missing.append(po)

    ack_text = ", ".join(ack_labels) if ack_labels else None
    status_text = "; ".join(status_labels) if status_labels else None

    if not pos and not needs_upwardor:
        recon = ScheduleRecon(ack_text, None, "", None)
    elif missing:
        listed = f" ({', '.join(missing)})" if ack_labels else ""
        recon = ScheduleRecon(
            ack_text,
            _join_status(status_text, f"Not acknowledged{listed}"),
            FLAG_RED,
            f"{NOTE_RED}{listed}",
        )
    elif not pos:
        recon = ScheduleRecon(None, "No Upwardor PO, no ack", FLAG_MISSED, NOTE_MISSED)
    elif _statuses_are_cancelled(status_text):
        recon = ScheduleRecon(ack_text, status_text, FLAG_CHECK, NOTE_CANCELLED)
    else:
        recon = ScheduleRecon(ack_text, status_text or "Acknowledged", FLAG_OK, None)

    if open_orders_by_sord is not None:
        _apply_open_order_list(recon, open_orders_by_sord)
    return recon


def index_open_orders(open_orders: Optional[Sequence[dict]]) -> Optional[Dict[str, dict]]:
    """None stays None (export not loaded). A sequence becomes {S-ORD: row}."""
    if open_orders is None:
        return None
    indexed: Dict[str, dict] = {}
    for row in open_orders:
        sord = str(row.get("sord") or "").strip().upper()
        if sord:
            indexed[sord] = row
    return indexed


def parse_upwardor_open_order_export(content: bytes) -> List[dict]:
    """Group an Upw Sales Order Master workbook into one dict per S-ORD.

    Raises ValueError when the export header is missing. A header with no
    data rows returns an empty list — that is a real empty open-order list,
    not a parse failure. Callers must not treat a parse error as "no export"
    or every acknowledged PO would be flagged as missing from the list.
    """
    if not content:
        raise ValueError("Upwardor open-order export is empty")
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        header_map = None
        for _ in range(15):
            try:
                raw = next(rows)
            except StopIteration:
                break
            header_map = _match_export_header(raw)
            if header_map is not None:
                break
        if header_map is None:
            raise ValueError(
                "Upwardor open-order export is missing a header row "
                "with 'Sales Order No.'"
            )
        missing = [name for name in _EXPORT_REQUIRED if name not in header_map]
        if missing:
            raise ValueError(
                "Upwardor open-order export is missing column(s): "
                + ", ".join(missing)
            )
        grouped: Dict[str, dict] = {}
        order: List[str] = []
        for raw in rows:
            if not raw or header_map[_EXPORT_SALES_ORDER] >= len(raw):
                continue
            sord_raw = raw[header_map[_EXPORT_SALES_ORDER]]
            sords = sord_numbers(sord_raw) or (
                [str(sord_raw).strip().upper()] if sord_raw else []
            )
            if not sords:
                continue
            sord = sords[0]
            qty_ordered = _float(_cell(raw, header_map, _EXPORT_QTY_ORDERED))
            qty_remaining = _float(_cell(raw, header_map, _EXPORT_QTY_REMAINING))
            qty_to_mfg = _float(_cell(raw, header_map, _EXPORT_QTY_TO_MFG))
            bulk = str(_cell(raw, header_map, _EXPORT_BULK) or "").strip()
            part = str(_cell(raw, header_map, _EXPORT_PART) or "").strip()
            description = str(_cell(raw, header_map, _EXPORT_DESCRIPTION) or "").strip()
            customer = str(_cell(raw, header_map, _EXPORT_CUSTOMER) or "").strip()
            bucket = grouped.get(sord)
            if bucket is None:
                bucket = {
                    "sord": sord,
                    "customer": customer,
                    "line_count": 0,
                    "sections_ordered": 0.0,
                    "sections_remaining": 0.0,
                    "sections_to_mfg": 0.0,
                    "lines": [],
                }
                grouped[sord] = bucket
                order.append(sord)
            elif customer and not bucket["customer"]:
                bucket["customer"] = customer
            bucket["line_count"] += 1
            if bulk.upper() == "SEC":
                bucket["sections_ordered"] += qty_ordered
                bucket["sections_remaining"] += qty_remaining
                bucket["sections_to_mfg"] += qty_to_mfg
            bucket["lines"].append({
                "part_no": part,
                "description": description,
                "qty_ordered": qty_ordered,
                "qty_remaining": qty_remaining,
                "qty_to_mfg": qty_to_mfg,
                "bulk": bulk,
            })
        return [grouped[sord] for sord in order]
    finally:
        wb.close()


def build_reconciliation_report(
    po_lines_by_so: Optional[Dict[str, List[dict]]],
    so_customer_map: Optional[Dict[str, str]],
    vendor_acks: Optional[Sequence[dict]],
    open_orders: Optional[Sequence[dict]] = None,
    unlinked_pos: Optional[Sequence[dict]] = None,
) -> ReconciliationReport:
    """Compare our POs (and any loaded Upwardor open orders) for the sheet.

    Without an export, the sheet still lists every Upwardor-relevant PO linked
    to an open SO, plus stock POs that reference no SO, and whether each has
    an acknowledgement. With an export, S-ORDs are matched through that
    acknowledgement. Part-level differences (suffixes, glass kits on another
    vendor) are left for a follow-up — `lines` on each parsed S-ORD are there
    for that.
    """
    acks = list(vendor_acks or [])
    groups = _collect_po_groups(po_lines_by_so or {}, so_customer_map or {}, acks, unlinked_pos or [])
    export_index = index_open_orders(open_orders)
    export_loaded = export_index is not None
    export_count = len(export_index or {})

    if export_loaded:
        source_note = (
            f"Compared with an Upwardor open-order export ({export_count} S-ORDs). "
            "Each S-ORD is matched to an OpenDC PO through the vendor acknowledgement. "
            "Part-level differences are not compared yet."
        )
    else:
        source_note = (
            "Upwardor open-order export is not loaded. This sheet lists OpenDC "
            "purchase orders and whether an Upwardor acknowledgement (S-ORD) is on file. "
            "Email ingest of the Upw Sales Order Master export is a follow-up — pass "
            "parse_upwardor_open_order_export() rows as upwardor_open_orders to compare."
        )

    open_order_rows: List[dict] = []
    if export_index is not None:
        ack_to_po = _ack_to_po(acks)
        for sord in sorted(export_index, key=_sord_sort):
            agg = export_index[sord]
            po = ack_to_po.get(sord)
            group = groups.get(po) if po else None
            open_order_rows.append(_open_order_row(sord, agg, po, group))

    our_rows = []
    listed = set(export_index or {})
    for po in sorted(groups, key=_po_sort):
        group = groups[po]
        sords = sord_numbers(group.get("vendor_ack"))
        if export_index is None:
            flag, note = _po_flag_without_export(group)
        elif not sords:
            flag, note = FLAG_RED, NOTE_RED
        elif any(sord not in listed for sord in sords):
            flag, note = FLAG_CHECK, NOTE_NOT_ON_LIST
        else:
            continue  # shown on the open-order table
        our_rows.append({
            "po_number": po,
            "vendor_ack": group.get("vendor_ack"),
            "so_number": ", ".join(group["so_numbers"]),
            "customer": ", ".join(c for c in group["customers"] if c),
            "po_status": group.get("po_status") or "",
            "flag": flag,
            "note": note,
        })
    our_rows.sort(key=lambda row: (_flag_rank(row["flag"]), _po_sort(row["po_number"])))

    return ReconciliationReport(
        export_loaded=export_loaded,
        source_note=source_note,
        open_order_rows=open_order_rows,
        our_po_rows=our_rows,
        export_count=export_count,
    )


def add_reconciliation_sheet(wb, report: ReconciliationReport) -> None:
    """Replace the Upwardor Reconciliation sheet. Nothing on it is hand-edited."""
    if RECON_SHEET_NAME in wb.sheetnames:
        del wb[RECON_SHEET_NAME]
    ws = wb.create_sheet(RECON_SHEET_NAME)
    ws.cell(row=1, column=1, value=report.source_note)
    ws.cell(row=1, column=1).font = _NOTE_FONT
    ws.cell(row=1, column=1).alignment = Alignment(wrap_text=True, vertical="center")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(OPEN_ORDER_HEADERS))
    ws.row_dimensions[1].height = 36
    ws.freeze_panes = "A3"

    row = 3
    if report.export_loaded:
        row = _write_section(ws, row, "Upwardor open orders", OPEN_ORDER_HEADERS, [
            [
                r["sord"], r["po_number"], r["match_source"], r["so_number"], r["customer"],
                r["line_count"], _display_num(r["sections_ordered"]),
                _display_num(r["sections_remaining"]), _display_num(r["sections_to_mfg"]),
                _display_num(r["bc_sections"]) if r["bc_sections"] is not None else None,
                r["result"], r["note"],
            ]
            for r in report.open_order_rows
        ], result_col=11)
        row += 1

    title = (
        "OpenDC POs with no acknowledgement, or not on the Upwardor open-order list"
        if report.export_loaded else
        "OpenDC purchase orders"
    )
    _write_section(ws, row, title, OUR_PO_HEADERS, [
        [r["po_number"], r["vendor_ack"], r["so_number"], r["customer"], r["po_status"], r["flag"], r["note"]]
        for r in report.our_po_rows
    ], result_col=6)

    widths = [22, 24, 28, 18, 28, 18, 20, 20, 18, 16, 36, 64]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


# ── internals ───────────────────────────────────────────────────────────

def _join_status(current: Optional[str], extra: str) -> str:
    if current and extra:
        return f"{current}; {extra}"
    return extra or current or ""


def _statuses_are_cancelled(status: Optional[str]) -> bool:
    if not status:
        return False
    parts = [part.strip() for part in status.replace(";", ",").split(",") if part.strip()]
    return bool(parts) and all(part.lower().startswith("cancelled") for part in parts)


def _apply_open_order_list(recon: ScheduleRecon, open_orders_by_sord: Dict[str, dict]) -> None:
    sords = sord_numbers(recon.ack_numbers)
    if not sords or recon.flag == FLAG_RED:
        return
    present = []
    missing = []
    for sord in sords:
        agg = open_orders_by_sord.get(sord)
        if not agg:
            missing.append(sord)
            continue
        blurb = _open_blurb(agg)
        present.append(f"{sord}: {blurb}" if len(sords) > 1 else blurb)
    if present:
        recon.upwardor_status = _join_status(recon.upwardor_status, "; ".join(present))
    if missing and recon.flag in (FLAG_OK, ""):
        recon.flag = FLAG_CHECK
        recon.note = _join_note(recon.note, f"{NOTE_NOT_ON_LIST} ({', '.join(missing)})")
        if not present:
            recon.upwardor_status = "Not on Upwardor open-order list"
    elif missing and recon.flag == FLAG_CHECK:
        recon.note = _join_note(recon.note, f"{NOTE_NOT_ON_LIST} ({', '.join(missing)})")


def _open_blurb(agg: dict) -> str:
    return (
        "Open on Upwardor list — sections "
        f"{_num(agg.get('sections_remaining'))} of {_num(agg.get('sections_ordered'))} remaining, "
        f"{_num(agg.get('sections_to_mfg'))} to mfg"
    )


def _join_note(current: Optional[str], extra: str) -> str:
    if current and extra:
        return f"{current}; {extra}"
    return extra or current or ""


def _is_operator_item(item_no: Optional[str]) -> bool:
    from app.services.production_schedule_service import OPERATOR_PREFIXES
    return (item_no or "").upper().startswith(OPERATOR_PREFIXES)


def _is_upwardor_vendor(name: Optional[str]) -> bool:
    text = (name or "").strip().upper()
    if not text:
        return True
    return "UPWARDOR" in text or text == "UPW"


def _is_section_item(item_no: Optional[str]) -> bool:
    return (item_no or "").upper().startswith(_BC_SECTION_PREFIXES)


def _bc_section_qty(lines: Iterable[dict]) -> float:
    total = 0.0
    for line in lines:
        if _is_operator_item(line.get("lineObjectNumber")):
            continue
        if not _is_section_item(line.get("lineObjectNumber")):
            continue
        total += _float(line.get("quantity"))
    return total


def _collect_po_groups(po_lines_by_so, so_customer_map, vendor_acks, unlinked_pos) -> Dict[str, dict]:
    groups: Dict[str, dict] = {}

    def ensure(po: str) -> dict:
        group = groups.get(po)
        if group is None:
            cells = ack_cells_for_po(po, vendor_acks)
            group = {
                "so_numbers": [],
                "customers": [],
                "po_status": "",
                "vendor": "",
                "lines": [],
                "_seen_lines": set(),
                "vendor_ack": cells.get("vendor_ack"),
                "ack_status": cells.get("ack_status"),
            }
            groups[po] = group
        return group

    for so, lines in po_lines_by_so.items():
        if so not in so_customer_map:
            continue
        for line in lines:
            if _is_operator_item(line.get("lineObjectNumber")):
                continue
            po = str(line.get("_po_number") or "").strip()
            if not po:
                continue
            group = ensure(po)
            if so not in group["so_numbers"]:
                group["so_numbers"].append(so)
                group["customers"].append(so_customer_map.get(so, "") or "")
            group["po_status"] = group["po_status"] or (line.get("_po_status") or "")
            group["vendor"] = group["vendor"] or (line.get("_vendor") or "")
            key = (
                line.get("sequence"),
                line.get("lineObjectNumber"),
                line.get("quantity"),
                line.get("description"),
            )
            if key in group["_seen_lines"]:
                continue
            group["_seen_lines"].add(key)
            group["lines"].append(line)

    for po in unlinked_pos:
        number = str(po.get("po_number") or "").strip()
        if not number or number in groups:
            continue
        if not _is_upwardor_vendor(po.get("vendor")):
            continue
        group = ensure(number)
        group["po_status"] = po.get("po_status") or ""
        group["vendor"] = po.get("vendor") or ""
        for line in po.get("lines") or []:
            if _is_operator_item(line.get("lineObjectNumber")):
                continue
            group["lines"].append(line)
    return groups


def _ack_to_po(vendor_acks: Sequence[dict]) -> Dict[str, str]:
    """S-ORD → PO number. Suffixed acks (PO-000960(2)) keep the base PO."""
    from app.services.vendor_ack_parser import po_base
    index: Dict[str, str] = {}
    for ack in vendor_acks:
        for sord in sord_numbers(ack.get("vendor_order_no")):
            po = po_base(ack.get("our_po_number")) or str(ack.get("our_po_number") or "").strip()
            if po and sord not in index:
                index[sord] = po
    return index


def _open_order_row(sord: str, agg: dict, po: Optional[str], group: Optional[dict]) -> dict:
    upw_sections = _float(agg.get("sections_ordered"))
    if not po:
        result, note, bc_sections = _RESULT_NO_PO, "No vendor acknowledgement links this S-ORD to an OpenDC PO", None
        match_source = "no PO reference found"
        so_number = ""
        customer = agg.get("customer") or ""
    elif group is None:
        result = _RESULT_CLOSED
        note = "Upwardor still lists this order but the PO is not open on an open sales order"
        match_source = "vendor ack"
        bc_sections = None
        so_number, customer = "", agg.get("customer") or ""
    else:
        bc_sections = _bc_section_qty(group["lines"])
        match_source = "vendor ack"
        so_number = ", ".join(group["so_numbers"])
        customer = ", ".join(c for c in group["customers"] if c) or (agg.get("customer") or "")
        if (upw_sections or bc_sections) and abs(upw_sections - bc_sections) > 0.001:
            result = _RESULT_QTY
            note = (
                f"Qty mismatch: BC {po} has {_num(bc_sections)} sections, "
                f"Upwardor {sord} has {_num(upw_sections)}"
            )
        else:
            result, note = _RESULT_MATCHED, None
    return {
        "sord": sord,
        "po_number": po,
        "match_source": match_source,
        "so_number": so_number,
        "customer": customer,
        "line_count": agg.get("line_count") or 0,
        "sections_ordered": upw_sections,
        "sections_remaining": _float(agg.get("sections_remaining")),
        "sections_to_mfg": _float(agg.get("sections_to_mfg")),
        "bc_sections": bc_sections,
        "result": result,
        "note": note,
    }


def _po_flag_without_export(group: dict):
    if group.get("vendor_ack"):
        if _statuses_are_cancelled(group.get("ack_status")):
            return FLAG_CHECK, NOTE_CANCELLED
        return FLAG_OK, None
    return FLAG_RED, NOTE_RED


def _flag_rank(flag: str) -> int:
    return {FLAG_RED: 0, FLAG_MISSED: 1, FLAG_CHECK: 2, FLAG_OK: 3}.get(flag, 4)


def _sord_sort(sord: str):
    digits = re.sub(r"\D", "", sord or "")
    return int(digits) if digits else 0


def _po_sort(po: str):
    return _sord_sort(po or "")


def _norm_header(value) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip()).lower()


def _match_export_header(raw) -> Optional[Dict[str, int]]:
    if not raw:
        return None
    mapping = {_norm_header(value): idx for idx, value in enumerate(raw) if value}
    if _EXPORT_SALES_ORDER not in mapping:
        return None
    return mapping


def _cell(raw, header_map, name):
    idx = header_map.get(name)
    if idx is None or idx >= len(raw):
        return None
    return raw[idx]


def _float(value) -> float:
    if value is None or value == "":
        return 0.0
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _num(value) -> str:
    number = _float(value)
    if number.is_integer():
        return str(int(number))
    return f"{number:g}"


def _display_num(value):
    if value is None:
        return None
    number = _float(value)
    return int(number) if number.is_integer() else number


def _write_section(ws, start_row: int, title: str, headers: List[str], data: List[list], result_col: int) -> int:
    title_cell = ws.cell(row=start_row, column=1, value=title)
    title_cell.fill = _SECTION_FILL
    title_cell.font = _SECTION_FONT
    header_row = start_row + 1
    for col, name in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=name)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for offset, values in enumerate(data):
        excel_row = header_row + 1 + offset
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=excel_row, column=col, value=value)
            cell.alignment = Alignment(vertical="center", wrap_text=(col == len(headers)))
        _paint_result(ws, excel_row, result_col, len(headers), values[result_col - 1] if values else None)
    return header_row + 1 + len(data)


def _paint_result(ws, row: int, result_col: int, width: int, result) -> None:
    text = str(result or "")
    if text in (FLAG_RED,):
        for col in range(1, width + 1):
            ws.cell(row=row, column=col).fill = _RED_FILL
            ws.cell(row=row, column=col).font = _RED_FONT
        return
    cell = ws.cell(row=row, column=result_col)
    if text in (FLAG_CHECK, _RESULT_QTY, _RESULT_CLOSED):
        cell.fill = _AMBER_FILL
        cell.font = _AMBER_FONT
    elif text == _RESULT_NO_PO:
        cell.fill = _ORANGE_FILL
        cell.font = _ORANGE_FONT
